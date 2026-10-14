from __future__ import annotations

from datetime import UTC, date, datetime
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.config import Settings
from app.database import Database
from app.models import NotificationRule, SheetSource
from app.services.notifications import build_test_notification_message
from app.services.sync_service import SyncService
from app.services.test_series import parse_test_workbook

TEST_SHEET_ID = "1ZYbITcTnAAbrLOyQ3pzaykEAmOt9Q1EH6gkImY33JXQ"


def source() -> SheetSource:
    return SheetSource(
        id="test-series-source",
        name="Test Series - 2026",
        url=f"https://docs.google.com/spreadsheets/d/{TEST_SHEET_ID}/edit",
        spreadsheet_id=TEST_SHEET_ID,
        enabled=True,
        created_at=datetime.now(UTC),
    )


def make_test_workbook(path: Path) -> None:
    workbook = Workbook()
    prarambh = workbook.active
    prarambh.title = "PRARAMBH 2026"
    prarambh.append([
        "S.No", "Dates", "Course", "Class", "Subject", "Chapter Number", "Chapter Name",
        "Made by", "No. Of Question", "Test Timing In minutes", "Test Marks", "Docs Link",
        "App Status", "Test Date", "Result Deshbord", "Video Solution Upodat",
    ])
    prarambh.append([
        1, date(2026, 8, 1), "Prarambh + Plus", "Class 11th", "Political Science",
        "Chapter 1", "Election and Representation", "Sahil", 20, 30, 80,
        "Election and Representation Test.docx", "Done", "05-08-2026 Done", "", "",
    ])
    prarambh["L2"].hyperlink = "https://docs.google.com/document/d/election/edit"

    nirmaan = workbook.create_sheet("Nirman 2026")
    nirmaan.append([
        "Dates", "Course", "Class", "Subject", "Chapter Number", "Chapter Name", "Made by",
        "No. Of Question", "Test Timing", "Test Marks", "Docs Link", "App Status", "Date",
    ])
    nirmaan.append([
        date(2026, 5, 27), "Nirman", "Class 8th", "Hindi", "Chapter 5", "कबीर के दोहे",
        "Kavita", 20, 45, 20, "कबीर के दोहे.docx", "Pending", "05-08-2026 Done",
    ])
    nirmaan["K2"].hyperlink = "https://docs.google.com/document/d/kabir/edit"
    workbook.save(path)


def test_structured_parser_finds_kabir_and_status_fields(tmp_path: Path) -> None:
    path = tmp_path / "tests.xlsx"
    make_test_workbook(path)
    records, content = parse_test_workbook(path, source_id="source", workbook_name="Test Series - 2026")
    assert len(records) == 2
    kabir = next(item for item in records if item.topic == "कबीर के दोहे")
    assert kabir.sheet_name == "Nirman 2026"
    assert kabir.app_status == "Pending"
    searchable = next(item for item in content if "कबीर" in item.title)
    assert searchable.record_kind == "test"
    assert searchable.status == "Pending"


def test_sync_baseline_then_detects_new_test(tmp_path: Path, database: Database) -> None:
    path = tmp_path / "tests.xlsx"
    make_test_workbook(path)
    item = source()
    database.upsert_source(item)
    service = SyncService(database, Settings(database_url=f"sqlite:///{database.path}"))
    first = service.sync_local_workbook(item, path, live_data=False)
    assert first.test_record_count == 2
    assert first.new_test_ids == []

    workbook = load_workbook(path)
    sheet = workbook["Nirman 2026"]
    sheet.append([
        date(2026, 8, 6), "Nirman", "Class 8th", "Sanskrit", "Chapter 6", "सुभाषितानि",
        "Vivek", 20, 45, 20, "Sanskrit Test.docx", "", "", 
    ])
    sheet["K3"].hyperlink = "https://docs.google.com/document/d/sanskrit/edit"
    workbook.save(path)
    second = service.sync_local_workbook(item, path, live_data=False)
    assert second.test_record_count == 3
    assert len(second.new_test_ids) == 1


def test_monitor_explains_remaining_work(
    tmp_path: Path, database: Database, client: TestClient
) -> None:
    path = tmp_path / "tests.xlsx"
    make_test_workbook(path)
    item = source()
    database.upsert_source(item)
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        item, path, live_data=False
    )
    response = client.get("/api/v1/test-monitor", params={"source_id": item.id})
    assert response.status_code == 200
    body = response.json()
    assert body["summary"]["total_tests"] == 2
    assert body["summary"]["app_status"]["pending"] == 1
    assert body["summary"]["result_dashboard"]["pending"] == 1
    assert body["summary"]["video_solution"]["pending"] == 1
    assert body["pending_app"][0]["topic"] == "कबीर के दोहे"


def test_one_click_presets_create_four_clear_rules(database: Database, client: TestClient) -> None:
    item = source()
    database.upsert_source(item)
    response = client.post(
        "/api/v1/notification-presets/test-monitor",
        json={
            "channel": "whatsapp_personal",
            "recipient": "+919876543210",
            "source_id": item.id,
            "app_status_time": "17:00",
            "result_time": "17:00",
            "video_time": "17:00",
        },
    )
    assert response.status_code == 200
    rules = response.json()
    assert {rule["event_type"] for rule in rules} == {
        "new_test_added", "test_app_status", "test_result_dashboard", "test_video_solution"
    }
    by_event = {rule["event_type"]: rule for rule in rules}
    assert by_event["test_app_status"]["weekday"] == 5
    assert by_event["test_result_dashboard"]["weekday"] == 0
    assert by_event["test_video_solution"]["weekday"] == 1
    assert by_event["new_test_added"]["send_time"] is None


def test_status_notification_lists_remaining_tests(tmp_path: Path, database: Database) -> None:
    path = tmp_path / "tests.xlsx"
    make_test_workbook(path)
    item = source()
    database.upsert_source(item)
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        item, path, live_data=False
    )
    rule = NotificationRule(
        id="result-rule",
        name="Monday Result Dashboard",
        event_type="test_result_dashboard",
        send_time="17:00",
        weekday=0,
        channels=["whatsapp_personal"],
        recipients=["+919876543210"],
        source_ids=[item.id],
        sheet_names=["PRARAMBH 2026"],
        created_at=datetime.now(UTC),
    )
    message = build_test_notification_message(database, rule, is_test=True)
    assert "Monday Result Dashboard report" in message
    assert "Remaining: 1" in message
    assert "Election and Representation" in message
