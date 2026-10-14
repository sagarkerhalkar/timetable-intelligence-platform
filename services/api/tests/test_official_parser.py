from __future__ import annotations

from datetime import date, datetime, time, timedelta

from openpyxl import Workbook

from app.timezones import INDIA_TIMEZONE

from app.services.official_parser import (
    OFFICIAL_SOURCES,
    canonical_subject,
    canonical_subject_group,
    parse_official_timetable,
    source_kind,
    source_stream,
    subjects_from_topic,
)


def fresh_workbook(title: str) -> tuple[Workbook, object]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    return workbook, sheet


def parse(workbook: Workbook, spreadsheet_id: str, name: str):
    return parse_official_timetable(
        workbook,
        spreadsheet_id=spreadsheet_id,
        source_id="source-1",
        workbook_name=name,
    )


def current_week_tuesday() -> date:
    today = datetime.now(INDIA_TIMEZONE).date()
    monday = today - timedelta(days=today.weekday())
    return monday + timedelta(days=1)


def previous_week_tuesday() -> date:
    return current_week_tuesday() - timedelta(days=7)


def test_official_source_helpers_and_subject_normalization() -> None:
    commerce_id = next(key for key, value in OFFICIAL_SOURCES.items() if value == "commerce")
    assert source_kind(commerce_id) == "commerce"
    assert source_stream("commerce") == "Commerce"
    assert source_kind("unknown") is None
    assert source_stream("unknown") is None
    assert canonical_subject("Accounts - 12") == "Accountancy"
    assert canonical_subject_group("Political Science") == "Political Science"
    assert canonical_subject_group("Applied Maths / Mathematics") == "Applied Mathematics"
    assert subjects_from_topic("Physics / Chemistry revision") == "Physics / Chemistry"
    assert parse(Workbook(), "unknown", "Unknown") is None


def test_commerce_weekly_planner_merges_prarambh_and_pro() -> None:
    workbook, sheet = fresh_workbook("Weekly Planner")
    sheet["A1"] = "Commerce Class 12 Prarambh"
    sheet.append(["Date", "Day", "Subject", "6:30 PM - 7:45 PM"])
    sheet.append([date(2026, 8, 4), "Tuesday", "Economics", "National Income"])
    sheet["A5"] = "Commerce Class 12 Prarambh Pro"
    sheet["A6"] = "Date"
    sheet["B6"] = "Day"
    sheet["C6"] = "Subject"
    sheet["D6"] = "6:30 PM - 7:45 PM"
    sheet["A7"] = date(2026, 8, 4)
    sheet["B7"] = "Tuesday"
    sheet["C7"] = "Economics"
    sheet["D7"] = "National Income"

    items = parse(
        workbook,
        "1vEGv2nzGYlFWtHT8SmvzGIYeL_Yl9RJNRohM2VHY0Lg",
        "Master Sheet Commerce - 11 /12 (2026-27)",
    )
    assert items is not None
    assert len(items) == 1
    item = items[0]
    assert item.class_date == date(2026, 8, 4)
    assert item.start_time == time(18, 30)
    assert item.end_time == time(19, 45)
    assert item.class_name == "Class 12"
    assert item.subject == "Economics"
    assert item.topic == "National Income"
    assert item.batch == "Prarambh + Prarambh Pro"


def test_science_planner_uses_topic_and_subject_tab_evidence() -> None:
    workbook, planner = fresh_workbook("Weekly Planner")
    planner["A1"] = "Prarambh Science Class 11"
    planner["A2"] = "Date"
    planner["B2"] = "Day"
    planner["C2"] = "5:00 PM"
    planner["D2"] = "6:30 PM"
    planner["A3"] = date(2026, 8, 4)
    planner["B3"] = "Tuesday"
    planner["C3"] = "Physics: Units and Measurements"
    planner["D3"] = "Motion in a Straight Line"

    physics = workbook.create_sheet("Physics - 11")
    physics["A1"] = "Chapter"
    physics["A2"] = "Motion in a Straight Line"

    items = parse(
        workbook,
        "1j1yR4DRzqJwnkkfHmeWuv7vP-iDP0KZVahV4l9lcjEg",
        "Prarambh Science (2026-27) - 11 & 12th",
    )
    assert items is not None
    assert len(items) == 2
    by_topic = {item.topic: item for item in items}
    assert by_topic["Physics: Units and Measurements"].subject == "Physics"
    assert by_topic["Motion in a Straight Line"].subject == "Physics"
    assert all(item.stream == "Science" for item in items)
    assert all(item.class_name == "Class 11" for item in items)


def test_science_falls_back_to_bounded_timetable_schedule() -> None:
    workbook, planner = fresh_workbook("Weekly Planner")
    planner["A1"] = "Prarambh Science Class 12"
    planner["A2"] = "Date"
    planner["B2"] = "Day"
    planner["C2"] = "5:00 PM"
    planner["A3"] = current_week_tuesday()
    planner["B3"] = "Tuesday"
    planner["C3"] = "Electrostatics revision session"

    timetable = workbook.create_sheet("Time Table")
    timetable["A1"] = "Prarambh Science Class 12"
    timetable["A2"] = "Day"
    timetable["B2"] = "5:00 PM"
    timetable["A3"] = "Tuesday"
    timetable["B3"] = "Physics"
    timetable["E1"] = "Commerce Class 12"
    timetable["E2"] = "Day"
    timetable["F2"] = "5:00 PM"
    timetable["E3"] = "Tuesday"
    timetable["F3"] = "Accountancy"

    items = parse(
        workbook,
        "1j1yR4DRzqJwnkkfHmeWuv7vP-iDP0KZVahV4l9lcjEg",
        "Prarambh Science (2026-27) - 11 & 12th",
    )
    assert items is not None and len(items) == 1
    assert items[0].subject == "Physics"
    assert "Accountancy" not in (items[0].subject or "")


def test_humanities_uses_schedule_without_misclassifying_political_science() -> None:
    workbook, planner = fresh_workbook("Weekly Planner")
    planner["C1"] = "Weekly Planner Class 11 Humanities"
    planner["C2"] = "5:00 PM"
    planner["A3"] = current_week_tuesday()
    planner["B3"] = "Tuesday"
    planner["C3"] = "Election and Representation"

    timetable = workbook.create_sheet("Time Table master")
    timetable["A1"] = "Day"
    timetable["B2"] = "Class 11 Humanities"
    timetable["B3"] = "5:00 PM"
    timetable["A4"] = "Tuesday"
    timetable["B4"] = "Political Science"

    items = parse(
        workbook,
        "1M8mxrLzxZ9AixQ7fejbzIjhZGLgNzS7flMx6cXgiXHQ",
        "Prarambh Batch - Class 11 & 12 Humanities | 2026-27",
    )
    assert items is not None and len(items) == 1
    item = items[0]
    assert item.stream == "Humanities"
    assert item.subject == "Political Science"
    assert item.topic == "Election and Representation"


def test_nirmaan_pairs_timetable_subject_with_planner_topic() -> None:
    workbook, sheet = fresh_workbook("Weekly Planner ")
    sheet["A1"] = "Nirmaan 2.0 Time Table"
    sheet["A2"] = "Date"
    sheet["B2"] = "Day"
    sheet["C2"] = "5:00 PM"
    sheet["A3"] = date(2026, 8, 4)
    sheet["B3"] = "Tuesday"
    sheet["C3"] = "Science"

    sheet["A6"] = "Nirmaan 2.0 Planner"
    sheet["A7"] = "Date"
    sheet["B7"] = "Day"
    sheet["C7"] = "5:00 PM"
    sheet["A8"] = date(2026, 8, 4)
    sheet["B8"] = "Tuesday"
    sheet["C8"] = "Crop Production and Management"

    items = parse(
        workbook,
        "1a7qJpddbGL3fnR6RBmsrFxqjjVLB_g9Ryz7cQ1i0W8U",
        "Nirmaan 2026-27 l Class 8 & 7",
    )
    assert items is not None and len(items) == 1
    item = items[0]
    assert item.stream == "Nirmaan"
    assert item.class_name == "Class 8"
    assert item.subject == "Science"
    assert item.topic == "Crop Production and Management"


def test_official_parsers_return_empty_when_required_tabs_are_missing() -> None:
    workbook = Workbook()
    assert parse(
        workbook,
        "1vEGv2nzGYlFWtHT8SmvzGIYeL_Yl9RJNRohM2VHY0Lg",
        "Commerce 2026-27",
    ) == []
    assert parse(
        workbook,
        "1j1yR4DRzqJwnkkfHmeWuv7vP-iDP0KZVahV4l9lcjEg",
        "Science 2026-27",
    ) == []
    assert parse(
        workbook,
        "1M8mxrLzxZ9AixQ7fejbzIjhZGLgNzS7flMx6cXgiXHQ",
        "Humanities 2026-27",
    ) == []
    assert parse(
        workbook,
        "1a7qJpddbGL3fnR6RBmsrFxqjjVLB_g9Ryz7cQ1i0W8U",
        "Nirmaan 2026-27",
    ) == []


def test_nirmaan_parses_every_ncert_variant_without_overwriting() -> None:
    workbook, sheet = fresh_workbook("Weekly Planner ")

    blocks = [
        (1, "NEW NCERT", "NIRMAAN BATCH - Time Table", "NIRMAAN BATCH - Planner", "Hindi", "Hindi grammar"),
        (13, "OLD NCERT", "NIRMAAN BATCH - Time Table", "NIRMAAN BATCH - Planner", "Science", "Force and Pressure L4"),
        (25, "NEW NCERT 2.0", "NIRMAAN 2.O BATCH - Time Table", "NIRMAAN BATCH - Planner", "Mathematics", "A Story of Numbers L3"),
    ]
    for start, marker, timetable_title, planner_title, subject, topic in blocks:
        sheet.cell(start, 1, marker)
        sheet.cell(start + 1, 1, timetable_title)
        sheet.cell(start + 2, 1, "Date")
        sheet.cell(start + 2, 2, "Days")
        sheet.cell(start + 2, 3, "17:00")
        sheet.cell(start + 3, 1, "Tuesday")
        sheet.cell(start + 3, 2, "4 Aug")
        sheet.cell(start + 3, 3, subject)
        sheet.cell(start + 5, 1, planner_title)
        sheet.cell(start + 6, 1, "Date")
        sheet.cell(start + 6, 2, "Days")
        sheet.cell(start + 6, 3, "17:00")
        sheet.cell(start + 7, 1, "Tuesday")
        sheet.cell(start + 7, 2, "4 Aug")
        sheet.cell(start + 7, 3, topic)

    items = parse(
        workbook,
        "1a7qJpddbGL3fnR6RBmsrFxqjjVLB_g9Ryz7cQ1i0W8U",
        "Nirmaan 2026-27 l Class 8 & 7",
    )
    assert items is not None and len(items) == 3
    by_batch = {item.batch: item for item in items}
    assert set(by_batch) == {"Nirmaan New NCERT", "Nirmaan Old NCERT", "Nirmaan 2.0"}
    assert by_batch["Nirmaan New NCERT"].subject == "Hindi"
    assert by_batch["Nirmaan Old NCERT"].topic == "Force and Pressure L4"
    assert by_batch["Nirmaan 2.0"].subject == "Mathematics"
    assert all(item.class_date == date(2026, 8, 4) for item in items)
    assert all(item.start_time == time(17, 0) for item in items)
    assert all(item.source_coordinates[0].startswith("C") for item in items)
    assert all(item.source_coordinates[1].startswith("C") for item in items)


def test_nirmaan_old_ncert_merged_six_pm_heading_keeps_both_classes() -> None:
    workbook, sheet = fresh_workbook("Weekly Planner ")
    sheet["A1"] = "OLD NCERT"
    sheet["A2"] = "NIRMAAN BATCH - Time Table"
    sheet["A3"] = "Date"
    sheet["B3"] = "Days"
    sheet["C3"] = "18:00:00 PM"
    # D3 is intentionally blank because the live Google Sheet visually merges
    # the 6 PM heading across C:D.
    sheet["A4"] = "Monday"
    sheet["B4"] = "3 Aug"
    sheet["C4"] = "Science"
    sheet["D4"] = "Maths"

    sheet["A7"] = "NIRMAAN BATCH - Planner"
    sheet["A8"] = "Date"
    sheet["B8"] = "Days"
    sheet["C8"] = "18:00:00 PM"
    sheet["A9"] = "Monday"
    sheet["B9"] = "3 Aug"
    sheet["C9"] = "Force and Pressure L3"
    sheet["D9"] = "Algebraic Expressions and Identities L3"

    items = parse(
        workbook,
        "1a7qJpddbGL3fnR6RBmsrFxqjjVLB_g9Ryz7cQ1i0W8U",
        "Nirmaan 2026-27 l Class 8 & 7",
    )

    assert items is not None and len(items) == 2
    assert {item.subject for item in items} == {"Science", "Mathematics"}
    assert {item.topic for item in items} == {
        "Force and Pressure L3",
        "Algebraic Expressions and Identities L3",
    }
    assert {item.start_time for item in items} == {time(18, 0)}
    assert {item.source_coordinates[0] for item in items} == {"C4", "D4"}
    assert {item.source_coordinates[1] for item in items} == {"C9", "D9"}


def test_commerce_uses_time_table_when_current_week_planner_is_missing() -> None:
    from app.services.official_parser import _current_week_dates

    workbook, planner = fresh_workbook("Weekly Planner")
    planner["A1"] = "Class 11th Commerce"
    planner["A2"] = "Date"
    planner["B2"] = "Day"
    planner["C2"] = "Subject"
    planner["D2"] = "17:00"
    previous_sunday = _current_week_dates()[0] - date.resolution
    planner["A3"] = previous_sunday
    planner["B3"] = "Sunday"
    planner["C3"] = "Economics"
    planner["D3"] = "Previous week topic"

    timetable = workbook.create_sheet("Time TablePlanner")
    timetable["A1"] = "Class 11th Commerce"
    timetable["A2"] = "Days"
    timetable["B2"] = "5:00 - 6:30 PM"
    timetable["C2"] = "8:00 - 9:30 PM"
    timetable["A3"] = "Wednesday"
    timetable["B3"] = "Economics"
    timetable["C3"] = "Accounts"

    items = parse(
        workbook,
        "1vEGv2nzGYlFWtHT8SmvzGIYeL_Yl9RJNRohM2VHY0Lg",
        "Master Sheet Commerce - 11 /12 (2026-27)",
    )
    target = _current_week_dates()[2]
    current = [item for item in items or [] if item.class_date == target]
    assert {(item.start_time, item.subject) for item in current} == {
        (time(17, 0), "Economics"),
        (time(20, 0), "Accountancy"),
    }
    assert all(item.topic == "Weekly Planner not updated — topic pending" for item in current)


def test_science_splits_parallel_math_and_biology_lessons() -> None:
    workbook, planner = fresh_workbook("Weekly Planner")
    planner["A1"] = "Prarambh Science Class 12"
    planner["A2"] = "Date"
    planner["B2"] = "Day"
    planner["C2"] = "8:00 PM"
    planner["A3"] = date(2026, 8, 6)
    planner["B3"] = "Thursday"
    planner["C3"] = "Application of Derivatives (L5) - Maths / Evolution (L5) - Biology"

    items = parse(
        workbook,
        "1j1yR4DRzqJwnkkfHmeWuv7vP-iDP0KZVahV4l9lcjEg",
        "Prarambh Science (2026-27) - 11 & 12th",
    )
    assert items is not None and len(items) == 2
    assert {item.subject for item in items} == {"Mathematics", "Biology"}
    assert all(item.start_time == time(20, 0) for item in items)


def test_science_previous_week_planner_adds_current_week_recurring_fallback() -> None:
    workbook, planner = fresh_workbook("Weekly Planner")
    planner["A1"] = "Prarambh Science Class 12"
    planner["A2"] = "Date"
    planner["B2"] = "Day"
    planner["C2"] = "5:00 PM"
    planner["A3"] = previous_week_tuesday()
    planner["B3"] = "Tuesday"
    planner["C3"] = "Electrostatics revision session"

    timetable = workbook.create_sheet("Time Table")
    timetable["A1"] = "Prarambh Science Class 12"
    timetable["A2"] = "Day"
    timetable["B2"] = "5:00 PM"
    timetable["A3"] = "Tuesday"
    timetable["B3"] = "Physics"

    items = parse(
        workbook,
        "1j1yR4DRzqJwnkkfHmeWuv7vP-iDP0KZVahV4l9lcjEg",
        "Prarambh Science (2026-27) - 11 & 12th",
    )
    assert items is not None
    physics = [item for item in items if item.subject == "Physics"]
    assert len(physics) == 2
    dates = sorted(item.class_date for item in physics if item.class_date)
    assert dates == [previous_week_tuesday(), current_week_tuesday()]
    fallback = next(item for item in physics if item.class_date == current_week_tuesday())
    assert fallback.topic == "Weekly Planner not updated — topic pending"
