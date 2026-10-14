from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from app.models import ContentType, TabKind
from app.services.workbook_analyzer import (
    WorkbookAnalyzer,
    classify_content_type,
    detect_class,
    detect_stream,
    detect_subject,
)


def build_sample_workbook(path: Path) -> None:
    workbook = Workbook()
    timetable = workbook.active
    timetable.title = "Class 12 Commerce Time Table"
    timetable.merge_cells("A1:H1")
    timetable["A1"] = "Class 12 Commerce Weekly Time Table"
    timetable.append(["Date", "Time", "Class", "Subject", "Topic", "Teacher", "Video", "PDF"])
    timetable.append(
        [
            date(2026, 8, 4),
            "6:30 PM - 7:30 PM",
            "Class 12",
            "Economics",
            "National Income",
            "Asha Ma'am",
            "https://youtu.be/example",
            "https://drive.google.com/file/d/example/view?usp=sharing PDF",
        ]
    )
    planner = workbook.create_sheet("Weekly Planner")
    planner.append(["Class", "Subject", "Chapter", "Link"])
    planner.append(["11th Science", "Political Science", "Election and Representation", "https://example.com/notes.pdf"])
    workbook.save(path)


def test_detect_tab_kind() -> None:
    kind, confidence = WorkbookAnalyzer.detect_tab_kind(
        "Weekly Planner",
        Counter({"topic": 3, "subject": 2}),
        date_count=1,
        time_count=0,
        link_count=2,
    )
    assert kind in {TabKind.WEEKLY_PLANNER, TabKind.MIXED}
    assert confidence >= 0.5


def test_analyze_workbook(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    build_sample_workbook(path)
    profile, snapshots, timetable, content = WorkbookAnalyzer().analyze(
        path, "source-1", "Sample Workbook"
    )
    assert len(profile.tabs) == 2
    assert profile.total_non_empty_cells > 10
    assert any(tab.merged_ranges for tab in profile.tabs)
    assert snapshots
    assert timetable
    entry = next(item for item in timetable if item.subject == "Economics")
    assert entry.class_date == date(2026, 8, 4)
    assert entry.start_time and entry.start_time.hour == 18
    assert entry.class_name == "Class 12"
    assert any(item.content_type == ContentType.VIDEO for item in content)
    assert any(item.content_type == ContentType.PDF for item in content)


def test_subject_stream_and_class_detection() -> None:
    assert detect_subject("Election and Representation - Pol.Science") == "Political Science"
    assert detect_subject("Alcohol, Phenol & Ether - Chemistry") == "Chemistry"
    assert detect_stream("Class 12 Accountancy Commerce") == "Commerce"
    assert detect_class("Prarambh Pro 11th Science") == "Class 11"


def test_content_classification() -> None:
    assert classify_content_type("https://example.com/file.pdf") == ContentType.PDF
    assert classify_content_type("https://youtu.be/xyz") == ContentType.VIDEO
    assert (
        classify_content_type("https://docs.google.com/spreadsheets/d/abc")
        == ContentType.SPREADSHEET
    )


def test_sparse_scan_does_not_expand_large_empty_grid() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Date"
    sheet["Z5000"] = "Last real value"
    assert sheet.max_row == 5000
    rows = WorkbookAnalyzer._sparse_rows(sheet)
    assert [row[0].row for row in rows] == [1, 5000]
    assert sum(len(row) for row in rows) == 2


def test_content_rows_with_chapter_and_lecture_names_are_not_headers(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hindi - Lali Mam"
    sheet.append(["Chapter Name", "Lecture Name", "Video Link", "PDF Link"])
    sheet.append([
        "Chapter 1 Grammar",
        "Chapter 1 Lecture 1.mp4",
        "https://drive.google.com/file/d/video-one/view",
        "https://drive.google.com/file/d/pdf-one/view?name=notes.pdf",
    ])
    path = tmp_path / "content.xlsx"
    workbook.save(path)

    _, _, _, content = WorkbookAnalyzer().analyze(path, "source-1", "Nirmaan")
    urls = {item.url for item in content}
    assert "https://drive.google.com/file/d/video-one/view" in urls
    assert "https://drive.google.com/file/d/pdf-one/view?name=notes.pdf" in urls
    assert len(content) == 2


def test_same_url_in_multiple_cells_keeps_each_sheet_context(tmp_path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Maths"
    second = workbook.create_sheet("Science")
    for sheet in (first, second):
        sheet.append(["Topic", "Video Link"])
        sheet.append([sheet.title + " lesson", "https://example.com/shared-video.mp4"])
    path = tmp_path / "shared.xlsx"
    workbook.save(path)

    _, _, _, content = WorkbookAnalyzer().analyze(path, "source-1", "Nirmaan")
    shared = [item for item in content if item.url == "https://example.com/shared-video.mp4"]
    assert len(shared) == 2
    assert {item.sheet_name for item in shared} == {"Maths", "Science"}


def test_content_handles_common_sheet_typos_and_hyphenated_class() -> None:
    assert classify_content_type(
        "https://drive.google.com/file/d/example/view",
        "Vidoe Drive Link",
    ) == ContentType.VIDEO
    assert classify_content_type(
        "https://www.canva.com/design/example/edit",
        "PPT",
    ) == ContentType.PRESENTATION
    assert detect_class("Class-12th Swahaa Series") == "Class 12"
