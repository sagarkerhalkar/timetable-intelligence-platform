from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.config import Settings
from app.database import Database
from app.models import SheetSource
from app.services.sync_service import SyncService
from app.timezones import INDIA_TIMEZONE


def make_workbook(
    path: Path,
    topic: str = "National Income",
    class_date: date = date(2026, 8, 4),
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Class 12 Commerce Time Table"
    sheet.append(["Date", "Time", "Class", "Subject", "Topic", "PDF"])
    sheet.append(
        [
            class_date,
            "6:30 PM - 7:30 PM",
            "Class 12",
            "Economics",
            topic,
            "https://example.com/economics.pdf",
        ]
    )
    workbook.save(path)


def source() -> SheetSource:
    return SheetSource(
        id="source-1",
        name="Commerce",
        url="https://docs.google.com/spreadsheets/d/abc/edit",
        spreadsheet_id="abc",
        enabled=True,
        created_at=datetime.now(UTC),
    )


def test_local_sync_and_change_history(tmp_path: Path, database: Database) -> None:
    workbook_path = tmp_path / "commerce.xlsx"
    make_workbook(workbook_path)
    sheet_source = source()
    database.upsert_source(sheet_source)
    service = SyncService(database, Settings(database_url=f"sqlite:///{database.path}"))

    first = service.sync_local_workbook(sheet_source, workbook_path, live_data=False)
    assert first.status == "local_import"
    assert first.live_data is False
    assert first.timetable_count == 1
    assert first.content_count == 1
    assert database.list_timetable()[0].topic == "National Income"
    assert database.list_content()[0].content_type.value == "pdf"

    make_workbook(workbook_path, topic="Money and Banking")
    second = service.sync_local_workbook(sheet_source, workbook_path, live_data=False)
    assert second.change_count >= 1
    assert any(change.new_value == "Money and Banking" for change in database.list_changes())


def test_data_routes_and_exports(tmp_path: Path, database: Database, client: TestClient) -> None:
    workbook_path = tmp_path / "commerce.xlsx"
    make_workbook(workbook_path)
    sheet_source = source()
    database.upsert_source(sheet_source)
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        sheet_source, workbook_path, live_data=False
    )

    timetable = client.get(
        "/api/v1/timetable",
        params={"stream": "commerce", "class_name": "12", "subject": "economics"},
    )
    assert timetable.status_code == 200
    assert timetable.json()["total"] == 1

    content = client.get(
        "/api/v1/content",
        params={"query": "economics", "content_type": "pdf", "subject": "economics"},
    )
    assert content.status_code == 200
    assert content.json()["total"] == 1

    make_workbook(workbook_path, topic="Money and Banking")
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        sheet_source, workbook_path, live_data=False
    )
    changes = client.get(
        "/api/v1/changes",
        params={"source_id": "source-1", "change_type": "updated", "period": "all"},
    )
    assert changes.status_code == 200
    assert changes.json()["total"] > 0

    changes_csv = client.get(
        "/api/v1/exports/changes/csv",
        params={"source_id": "source-1", "period": "all"},
    )
    assert changes_csv.status_code == 200
    assert changes_csv.headers["content-type"].startswith("text/csv")
    assert b"Money and Banking" in changes_csv.content

    xlsx = client.get("/api/v1/exports/timetable/xlsx", params={"class_date": "2026-08-04"})
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    pdf = client.get("/api/v1/exports/timetable/pdf")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")


def test_missing_and_disabled_source_sync(client: TestClient, database: Database) -> None:
    missing = client.post("/api/v1/sources/not-found/sync")
    assert missing.status_code == 404

    disabled = source().model_copy(update={"enabled": False})
    database.upsert_source(disabled)
    response = client.post(f"/api/v1/sources/{disabled.id}/sync")
    assert response.status_code == 409


def test_out_of_range_page(client: TestClient, tmp_path: Path, database: Database) -> None:
    workbook_path = tmp_path / "commerce.xlsx"
    make_workbook(workbook_path)
    sheet_source = source()
    database.upsert_source(sheet_source)
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        sheet_source, workbook_path, live_data=False
    )
    response = client.get("/api/v1/timetable?page=99&page_size=1")
    assert response.status_code == 404


def test_background_sync_starts_without_holding_request(
    client: TestClient, database: Database, monkeypatch
) -> None:
    sheet_source = source()
    database.upsert_source(sheet_source)

    async def fake_sync(service: SyncService, item: SheetSource):
        service.database.update_source_sync_status(
            item.id, "success", None, datetime.now(UTC)
        )
        return None

    monkeypatch.setattr(SyncService, "sync_source", fake_sync)
    response = client.post(f"/api/v1/sources/{sheet_source.id}/sync/start")
    assert response.status_code == 202
    assert response.json()["status"] == "queued"
    current = client.get(f"/api/v1/sources/{sheet_source.id}")
    assert current.status_code == 200
    assert current.json()["last_sync_status"] == "success"


def test_timetable_today_and_upcoming_modes(
    tmp_path: Path, database: Database, client: TestClient
) -> None:
    local_today = datetime.now(INDIA_TIMEZONE).date()
    next_date = local_today + timedelta(days=1)
    workbook_path = tmp_path / "commerce.xlsx"
    make_workbook(workbook_path, class_date=next_date)
    sheet_source = source()
    database.upsert_source(sheet_source)
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        sheet_source, workbook_path, live_data=False
    )

    exact = client.get(
        "/api/v1/timetable",
        params={"class_date": next_date.isoformat(), "latest": False},
    )
    assert exact.status_code == 200
    assert exact.json()["total"] == 1

    upcoming = client.get("/api/v1/timetable", params={"upcoming": True, "latest": False})
    assert upcoming.status_code == 200
    assert upcoming.json()["total"] == 1
    assert upcoming.json()["items"][0]["class_date"] == next_date.isoformat()


def test_content_search_matches_exact_url_and_punctuation(
    tmp_path: Path, database: Database, client: TestClient
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Class 8 Science"
    sheet.append(["Class", "Subject", "Topic", "Video Link"])
    url = "https://drive.google.com/file/d/health-video/view?usp=sharing"
    sheet.append(["Class 8", "Science", "Health : The Ultimate Treasure | L2", url])
    path = tmp_path / "search.xlsx"
    workbook.save(path)
    sheet_source = source().model_copy(update={"name": "Nirmaan"})
    database.upsert_source(sheet_source)
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        sheet_source, path, live_data=False
    )

    by_url = client.get("/api/v1/content", params={"query": url})
    assert by_url.status_code == 200
    assert by_url.json()["total"] == 1

    by_words = client.get(
        "/api/v1/content",
        params={"query": "Health The Ultimate Treasure L2"},
    )
    assert by_words.status_code == 200
    assert by_words.json()["total"] == 1


def test_content_search_supports_single_words_fuzzy_and_multilingual(
    tmp_path: Path, database: Database, client: TestClient
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Class 8 Science"
    sheet.append(["Class", "Subject", "Topic", "Video Link"])
    sheet.append([
        "Class 8",
        "Science",
        "Light Reflection and Refraction | प्रकाश का परावर्तन",
        "https://example.com/light-reflection.mp4",
    ])
    sheet.append([
        "Class 8",
        "Sanskrit",
        "संस्कृत व्याकरण",
        "https://example.com/sanskrit-grammar.pdf",
    ])
    path = tmp_path / "multilingual-search.xlsx"
    workbook.save(path)
    sheet_source = source().model_copy(update={"name": "Nirmaan"})
    database.upsert_source(sheet_source)
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        sheet_source, path, live_data=False
    )

    for query in ("light", "प्रकाश", "scince", "संस्कृत", "sanskrit"):
        response = client.get("/api/v1/content", params={"query": query})
        assert response.status_code == 200
        assert response.json()["total"] >= 1, query


def test_timetable_date_range(
    tmp_path: Path, database: Database, client: TestClient
) -> None:
    source_item = source()
    database.upsert_source(source_item)
    service = SyncService(database, Settings(database_url=f"sqlite:///{database.path}"))
    workbook_path = tmp_path / "range.xlsx"
    make_workbook(workbook_path, class_date=date(2026, 8, 4))
    service.sync_local_workbook(source_item, workbook_path, live_data=False)

    response = client.get(
        "/api/v1/timetable",
        params={"date_from": "2026-08-03", "date_to": "2026-08-09"},
    )
    assert response.status_code == 200
    assert response.json()["total"] == 1


def test_content_search_reports_fast_index_metadata(
    tmp_path: Path, database: Database, client: TestClient
) -> None:
    workbook_path = tmp_path / "indexed-search.xlsx"
    make_workbook(workbook_path, topic="Light and Reflection")
    item = source()
    database.upsert_source(item)
    SyncService(database, Settings(database_url=f"sqlite:///{database.path}")).sync_local_workbook(
        item, workbook_path, live_data=False
    )
    response = client.get("/api/v1/content", params={"query": "light"})
    assert response.status_code == 200
    assert response.json()["total"] == 1
    assert response.json()["search_ms"] >= 1
    assert response.json()["normalized_query"] == "light"


def test_weekly_timetable_allows_large_page_size(client: TestClient) -> None:
    response = client.get(
        "/api/v1/timetable",
        params={"date_from": "2026-08-03", "date_to": "2026-08-09", "page_size": 1000},
    )
    assert response.status_code == 200
    assert response.json()["page_size"] == 1000
