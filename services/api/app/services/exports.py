from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..models import TimetableEntry

HEADERS = ["Date", "Day", "Time", "Stream", "Class", "Batch", "Subject", "Topic", "Teacher", "Source"]


def timetable_rows(entries: list[TimetableEntry]) -> list[list[str]]:
    rows: list[list[str]] = []
    for entry in entries:
        time_text = ""
        if entry.start_time:
            time_text = entry.start_time.strftime("%I:%M %p")
        if entry.end_time:
            time_text += f" - {entry.end_time.strftime('%I:%M %p')}"
        rows.append(
            [
                entry.class_date.isoformat() if entry.class_date else "",
                entry.day or "",
                time_text,
                entry.stream or "",
                entry.class_name or "",
                entry.batch or "",
                entry.subject or "",
                entry.topic or "",
                entry.teacher or "",
                f"{entry.workbook_name} / {entry.sheet_name} / row {entry.source_row}",
            ]
        )
    return rows


def build_timetable_xlsx(entries: list[TimetableEntry]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timetable"
    sheet.append(HEADERS)
    for row in timetable_rows(entries):
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="173B57")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    widths = [13, 12, 22, 16, 14, 18, 20, 36, 20, 45]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_timetable_pdf(entries: list[TimetableEntry], title: str = "Daily Timetable") -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=8 * mm,
        leftMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    data = [[Paragraph(f"<b>{header}</b>", styles["BodyText"]) for header in HEADERS[:-1]]]
    for row in timetable_rows(entries):
        data.append([Paragraph(value.replace("&", "&amp;"), styles["BodyText"]) for value in row[:-1]])
    table = Table(data, repeatRows=1, colWidths=[22*mm, 20*mm, 31*mm, 25*mm, 23*mm, 28*mm, 28*mm, 55*mm, 30*mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173B57")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8C4CE")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story = [Paragraph(title, styles["Title"]), Spacer(1, 5 * mm), table]
    document.build(story)
    return output.getvalue()
