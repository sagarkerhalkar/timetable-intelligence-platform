from __future__ import annotations

import re
from collections import Counter, defaultdict
from io import BytesIO
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from ..models import (
    CellSnapshot,
    ContentItem,
    ContentType,
    TabKind,
    TabProfile,
    TimetableEntry,
    WorkbookProfile,
)
from ..utils import find_urls, normalize_text, stable_hash
from .sheet_schema import profile_worksheet
from .official_parser import (
    canonical_subject,
    parse_official_timetable,
    source_kind,
    source_stream,
)

TIME_RE = re.compile(
    r"\b(?P<hour>0?[1-9]|1[0-2])[:.]?(?P<minute>[0-5][0-9])?\s*(?P<meridiem>a\.?m\.?|p\.?m\.?)\b",
    re.IGNORECASE,
)
TIME_24_RE = re.compile(r"\b(?P<hour>[01]?\d|2[0-3]):(?P<minute>[0-5]\d)\b")
DATE_RE = re.compile(
    r"\b(?:0?[1-9]|[12]\d|3[01])[-/.](?:0?[1-9]|1[0-2])[-/.](?:20)?\d{2}\b"
)
CLASS_RE = re.compile(r"\b(?:class|grade|std\.?|standard)?[\s-]*(7|8|9|10|11|12)(?:th|st|nd|rd)?\b", re.I)

SUBJECT_ALIASES: dict[str, tuple[str, ...]] = {
    "Accountancy": ("accountancy", "accounts"),
    "Business Studies": ("business studies", "bst"),
    "Economics": ("economics", "microeconomics", "macroeconomics", "indian economy"),
    "Applied Mathematics": ("applied mathematics", "applied maths"),
    "Mathematics": ("mathematics", "maths", "math", "गणित"),
    "Physics": ("physics", "भौतिक", "भौतिकी"),
    "Chemistry": ("chemistry", "organic chemistry", "inorganic chemistry", "physical chemistry", "रसायन", "रसायन विज्ञान"),
    "Biology": ("biology", "botany", "zoology", "जीव विज्ञान", "जीवविज्ञान"),
    "Political Science": ("political science", "pol. science", "pol science", "pol.science", "राजनीति विज्ञान"),
    "History": ("history", "इतिहास"),
    "Geography": ("geography", "geoegraphy", "geograpgy", "भूगोल"),
    "Psychology": ("psychology",),
    "Sociology": ("sociology",),
    "English": ("english", "अंग्रेजी", "अँग्रेज़ी"),
    "Hindi": ("hindi", "हिंदी", "हिन्दी"),
    "Sanskrit": ("sanskrit", "sanskirt", "संस्कृत", "संस्कृतम्"),
    "Computer Science": ("computer science", "informatics practices", "ip"),
    "Entrepreneurship": ("entrepreneurship",),
    "Physical Education": ("physical education",),
    "Science": ("science", "विज्ञान"),
    "Social Science": ("social science", "sst", "सामाजिक विज्ञान"),
}

STREAM_ALIASES: dict[str, tuple[str, ...]] = {
    "Commerce": ("commerce", "accountancy", "accounts", "business studies", "bst", "entrepreneurship"),
    "Science": ("science", "pcm", "pcb", "pcmb", "physics", "chemistry", "biology"),
    "Humanities": (
        "humanities",
        "arts",
        "political science",
        "history",
        "geography",
        "psychology",
        "sociology",
    ),
    "Nirmaan": ("nirmaan", "class 8", "grade 8", "8th"),
}

CONTENT_EXTENSIONS: dict[str, ContentType] = {
    ".pdf": ContentType.PDF,
    ".mp4": ContentType.VIDEO,
    ".mov": ContentType.VIDEO,
    ".webm": ContentType.VIDEO,
    ".mkv": ContentType.VIDEO,
    ".doc": ContentType.DOCUMENT,
    ".docx": ContentType.DOCUMENT,
    ".xls": ContentType.SPREADSHEET,
    ".xlsx": ContentType.SPREADSHEET,
    ".ppt": ContentType.PRESENTATION,
    ".pptx": ContentType.PRESENTATION,
    ".png": ContentType.IMAGE,
    ".jpg": ContentType.IMAGE,
    ".jpeg": ContentType.IMAGE,
}


class WorkbookAnalyzer:
    def analyze(
        self,
        workbook_path: Path,
        source_id: str,
        workbook_name: str | None = None,
        spreadsheet_id: str | None = None,
    ) -> tuple[WorkbookProfile, list[CellSnapshot], list[TimetableEntry], list[ContentItem]]:
        workbook_input = workbook_path if workbook_path.suffix.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"} else BytesIO(workbook_path.read_bytes())
        workbook = load_workbook(workbook_input, data_only=False, read_only=False)
        name = workbook_name or workbook_path.stem
        profiles: list[TabProfile] = []
        snapshots: list[CellSnapshot] = []
        timetable: list[TimetableEntry] = []
        content: list[ContentItem] = []
        warnings: list[str] = []
        official_kind = source_kind(spreadsheet_id)
        default_stream = source_stream(official_kind)

        for worksheet in workbook.worksheets:
            # Google Sheets often reports grids with 1,000+ rows even when only a small
            # percentage of cells contain data. Build one sparse row cache and reuse it
            # for profiling, change history, timetable extraction and content indexing.
            # This avoids repeatedly walking millions of empty cell coordinates.
            sparse_rows = self._sparse_rows(worksheet)
            profile = self.profile_tab(worksheet, sparse_rows)
            profiles.append(profile)
            snapshots.extend(self.snapshot_tab(worksheet, source_id, name, sparse_rows))
            if official_kind is None and profile.kind in {TabKind.TIMETABLE, TabKind.WEEKLY_PLANNER, TabKind.MIXED}:
                timetable.extend(self.extract_timetable(worksheet, source_id, name, sparse_rows))
            content.extend(
                self.extract_content(
                    worksheet, source_id, name, sparse_rows, default_stream=default_stream
                )
            )
            warnings.extend(f"{worksheet.title}: {warning}" for warning in profile.warnings)

        official_timetable = parse_official_timetable(
            workbook,
            spreadsheet_id=spreadsheet_id or "",
            source_id=source_id,
            workbook_name=name,
        )
        if official_timetable is not None:
            timetable = official_timetable
            warnings.append(
                f"Verified {official_kind} parser used for timetable and planner reconciliation"
            )

        profile = WorkbookProfile(
            source_id=source_id,
            workbook_name=name,
            analyzed_at=datetime.now().astimezone(),
            tabs=profiles,
            total_non_empty_cells=sum(item.non_empty_cells for item in profiles),
            total_links=sum(item.link_count for item in profiles),
            warnings=warnings,
        )
        return profile, snapshots, self._dedupe_timetable(timetable), self._dedupe_content(content)

    def profile_tab(
        self, worksheet: Worksheet, rows: Sequence[Sequence[Cell]] | None = None
    ) -> TabProfile:
        non_empty = 0
        links = 0
        date_count = 0
        time_count = 0
        row_density: Counter[int] = Counter()
        keywords: Counter[str] = Counter()

        sparse_rows = rows if rows is not None else self._sparse_rows(worksheet)
        row_lookup = {row[0].row: row for row in sparse_rows if row}
        for row in sparse_rows:
            for cell in row:
                value = cell.value
                text = normalize_text(value)
                hyperlink = cell.hyperlink.target if cell.hyperlink else None
                if not text and not hyperlink:
                    continue
                non_empty += 1
                row_density[cell.row] += 1
                lower = text.lower()
                links += int(bool(hyperlink)) + len(find_urls(text))
                date_count += int(self._looks_like_date(value, text))
                time_count += int(self._looks_like_time(value, text))
                for keyword in (
                    "time table",
                    "timetable",
                    "weekly planner",
                    "planner",
                    "subject",
                    "topic",
                    "chapter",
                    "class",
                    "batch",
                    "teacher",
                    "video",
                    "pdf",
                    "date",
                    "time",
                ):
                    if keyword in lower:
                        keywords[keyword] += 1

        kind, confidence = self.detect_tab_kind(
            worksheet.title,
            keywords,
            date_count=date_count,
            time_count=time_count,
            link_count=links,
        )
        header_candidates = [
            row_number
            for row_number, _ in row_density.most_common(8)
            if self._cells_have_header_words(row_lookup.get(row_number, ()))
        ][:5]
        schema = profile_worksheet(worksheet)
        warnings: list[str] = []
        if non_empty == 0:
            warnings.append("Tab is empty")
        if kind == TabKind.UNKNOWN and non_empty and not schema.search_ready:
            warnings.append("The app could not confidently understand the headings")
        if schema.header_row and not schema.search_ready and non_empty:
            warnings.append("Headings were found, but important search fields are missing")
        if worksheet.max_row > 5000:
            warnings.append("Large tab detected; sync will use streaming/chunking in production")

        return TabProfile(
            name=worksheet.title,
            kind=kind,
            max_row=worksheet.max_row,
            max_column=worksheet.max_column,
            non_empty_cells=non_empty,
            merged_ranges=[str(item) for item in worksheet.merged_cells.ranges],
            header_candidates=sorted(header_candidates),
            detected_keywords=sorted(keywords),
            link_count=links,
            date_count=date_count,
            time_count=time_count,
            confidence=confidence,
            warnings=warnings,
            hidden=worksheet.sheet_state != "visible",
            header_row=schema.header_row,
            detected_columns={field: schema.labels.get(field, field) for field in schema.searchable_fields},
            searchable_fields=schema.searchable_fields,
            missing_recommended_fields=schema.missing_recommended_fields,
            search_ready=schema.search_ready or kind in {TabKind.TIMETABLE, TabKind.WEEKLY_PLANNER, TabKind.MIXED},
            record_count=schema.record_count,
        )

    @staticmethod
    def detect_tab_kind(
        title: str,
        keywords: Counter[str] | dict[str, int],
        *,
        date_count: int,
        time_count: int,
        link_count: int,
    ) -> tuple[TabKind, float]:
        title_lower = title.lower()
        planner_score = 0
        timetable_score = 0
        content_score = 0
        if "weekly" in title_lower or "planner" in title_lower:
            planner_score += 6
        if "time table" in title_lower or "timetable" in title_lower or "schedule" in title_lower:
            timetable_score += 6
        if any(word in title_lower for word in ("content", "video", "pdf", "material", "lecture")):
            content_score += 5
        planner_score += int(keywords.get("weekly planner", 0)) * 3
        planner_score += int(keywords.get("planner", 0))
        planner_score += int(keywords.get("topic", 0)) + int(keywords.get("chapter", 0))
        timetable_score += int(keywords.get("time table", 0)) * 3
        timetable_score += int(keywords.get("timetable", 0)) * 3
        timetable_score += min(time_count, 8)
        timetable_score += min(date_count, 4)
        content_score += min(link_count, 10)
        content_score += int(keywords.get("video", 0)) + int(keywords.get("pdf", 0))

        scores = {
            TabKind.WEEKLY_PLANNER: planner_score,
            TabKind.TIMETABLE: timetable_score,
            TabKind.CONTENT_LIBRARY: content_score,
        }
        ordered = sorted(scores.items(), key=lambda item: item[1], reverse=True)
        top_kind, top_score = ordered[0]
        second_score = ordered[1][1]
        if top_score < 3:
            return TabKind.UNKNOWN, 0.25
        if top_score >= 5 and second_score >= 5 and abs(top_score - second_score) <= 3:
            return TabKind.MIXED, min(0.95, 0.55 + top_score / 30)
        return top_kind, min(0.98, 0.50 + top_score / 25)

    def snapshot_tab(
        self,
        worksheet: Worksheet,
        source_id: str,
        workbook_name: str,
        rows: Sequence[Sequence[Cell]] | None = None,
    ) -> list[CellSnapshot]:
        snapshots: list[CellSnapshot] = []
        for row in rows if rows is not None else self._sparse_rows(worksheet):
            for cell in row:
                value = cell.value
                text = normalize_text(value)
                hyperlink = cell.hyperlink.target if cell.hyperlink else None
                if not text and not hyperlink:
                    continue
                formula = text if isinstance(value, str) and value.startswith("=") else None
                snapshots.append(
                    CellSnapshot(
                        source_id=source_id,
                        workbook_name=workbook_name,
                        sheet_name=worksheet.title,
                        coordinate=cell.coordinate,
                        row=cell.row,
                        column=cell.column,
                        value=text,
                        formula=formula,
                        hyperlink=hyperlink,
                        style_key=self._style_key(cell),
                    )
                )
        return snapshots

    def extract_content(
        self,
        worksheet: Worksheet,
        source_id: str,
        workbook_name: str,
        rows: Sequence[Sequence[Cell]] | None = None,
        *,
        default_stream: str | None = None,
    ) -> list[ContentItem]:
        items: list[ContentItem] = []
        sparse_rows = rows if rows is not None else self._sparse_rows(worksheet)
        headers = self._best_header_map(worksheet, sparse_rows)
        title_context = self._context_from_text(worksheet.title)
        carried: dict[str, str | None] = {
            "chapter": None,
            "topic": None,
            "book": None,
            "teacher": None,
        }
        for row in sparse_rows:
            row_values = self._dense_values(row)
            # Repeated header rows are common inside long subject tabs. Never
            # carry words such as "Topic" or "Chapter Name" into subsequent
            # content records as real learning metadata.
            if self._looks_like_content_header(row_values):
                continue
            row_context = self._context_from_row(row_values, headers, worksheet.title)
            explicit: dict[str, str] = {}
            for index, value in enumerate(row_values, start=1):
                if not value:
                    continue
                header = headers.get(index, "").casefold()
                if "chapter" in header:
                    explicit["chapter"] = value
                elif "topic" in header:
                    explicit["topic"] = value
                elif "book" in header:
                    explicit["book"] = value
                elif (
                    "teacher" in header
                    or "faculty" in header
                    or "made by" in header
                    or "prepared by" in header
                ) and "uploaded" not in header:
                    explicit["teacher"] = value
            for key, value in explicit.items():
                carried[key] = value

            stream = default_stream or row_context["stream"] or title_context["stream"]
            class_name = row_context["class_name"] or title_context["class_name"]
            if default_stream == "Nirmaan" and class_name is None:
                class_name = "Class 8"
            subject = row_context["subject"] or title_context["subject"] or canonical_subject(worksheet.title)
            topic = explicit.get("topic") or carried.get("topic") or explicit.get("chapter") or carried.get("chapter")
            teacher = explicit.get("teacher") or carried.get("teacher")

            for cell in row:
                urls: list[str] = []
                if cell.hyperlink and cell.hyperlink.target:
                    urls.append(cell.hyperlink.target)
                urls.extend(find_urls(cell.value))
                for url in dict.fromkeys(urls):
                    label = normalize_text(cell.value)
                    header = headers.get(cell.column, "")
                    title = label if label and label != url else self._nearest_title(worksheet, cell)
                    content_type = classify_content_type(url, f"{header} {title}")
                    if _skip_non_learning_link(header, url):
                        continue
                    content_subject = (
                        subject
                        or canonical_subject(title)
                        or detect_subject(title)
                        or _subject_from_tab_topic(worksheet.title, title, topic)
                        or _subject_from_teacher(teacher, default_stream)
                    )
                    content_topic = topic or _clean_content_title(title)
                    # Preserve every sheet occurrence. The same Drive file can be
                    # intentionally listed under different classes, subjects or topics.
                    # Dedupe only the exact cell occurrence, not the URL workbook-wide.
                    fingerprint = stable_hash(
                        source_id, worksheet.title, cell.coordinate, _normalize_url(url)
                    )
                    confidence = 0.99 if cell.hyperlink else 0.88
                    warnings: list[str] = []
                    if not content_subject:
                        warnings.append("Subject was not detected from the tab, row, title or teacher")
                    items.append(
                        ContentItem(
                            id=stable_hash("content", source_id, worksheet.title, cell.coordinate, url),
                            source_id=source_id,
                            workbook_name=workbook_name,
                            sheet_name=worksheet.title,
                            coordinate=cell.coordinate,
                            source_row=cell.row,
                            url=url,
                            title=title or topic or "Untitled content",
                            content_type=content_type,
                            stream=stream,
                            class_name=class_name,
                            batch=row_context["batch"],
                            subject=content_subject,
                            topic=content_topic,
                            teacher=teacher,
                            search_text=" ".join(
                                value
                                for value in (
                                    worksheet.title,
                                    " ".join(row_values),
                                    title,
                                    topic,
                                    carried.get("chapter"),
                                    carried.get("topic"),
                                    teacher,
                                    content_subject,
                                    stream,
                                    class_name,
                                )
                                if value
                            ),
                            confidence=confidence,
                            warnings=warnings,
                            fingerprint=fingerprint,
                        )
                    )
        return items

    def extract_timetable(
        self,
        worksheet: Worksheet,
        source_id: str,
        workbook_name: str,
        rows: Sequence[Sequence[Cell]] | None = None,
    ) -> list[TimetableEntry]:
        entries: list[TimetableEntry] = []
        sparse_rows = rows if rows is not None else self._sparse_rows(worksheet)
        headers = self._best_header_map(worksheet, sparse_rows)
        sheet_context = self._context_from_text(worksheet.title)
        carried_date: date | None = None
        carried_class: str | None = sheet_context["class_name"]
        carried_stream: str | None = sheet_context["stream"]
        carried_batch: str | None = None

        for row in sparse_rows:
            cells = list(row)
            values = self._dense_values(cells)
            joined = " | ".join(value for value in values if value)
            if not joined:
                continue
            row_date = self._find_date(cells) or carried_date
            row_times = self._find_times(cells)
            row_context = self._context_from_row(values, headers, worksheet.title)
            carried_date = self._find_date(cells) or carried_date
            carried_class = row_context["class_name"] or carried_class
            carried_stream = row_context["stream"] or carried_stream
            carried_batch = row_context["batch"] or carried_batch

            if not row_times:
                continue
            start_time = row_times[0]
            end_time = row_times[1] if len(row_times) > 1 else None
            subject = row_context["subject"]
            topic = row_context["topic"]
            warnings: list[str] = []
            confidence = 0.55
            if row_date:
                confidence += 0.12
            else:
                warnings.append("Date was not confidently detected")
            if subject:
                confidence += 0.16
            else:
                warnings.append("Subject was not confidently detected")
            if topic:
                confidence += 0.08
            if carried_class:
                confidence += 0.05
            if subject and topic and normalize_text(subject).lower() == normalize_text(topic).lower():
                warnings.append("Subject-only/template row may require merging with planner data")
                confidence -= 0.18

            coordinates = [cell.coordinate for cell in cells if normalize_text(cell.value)]
            fingerprint = stable_hash(
                row_date,
                start_time,
                end_time,
                carried_stream,
                carried_class,
                carried_batch,
                subject,
                topic,
            )
            entries.append(
                TimetableEntry(
                    id=stable_hash("timetable", source_id, worksheet.title, row[0].row, fingerprint),
                    source_id=source_id,
                    workbook_name=workbook_name,
                    sheet_name=worksheet.title,
                    source_row=row[0].row,
                    source_coordinates=coordinates,
                    class_date=row_date,
                    day=row_date.strftime("%A") if row_date else None,
                    start_time=start_time,
                    end_time=end_time,
                    stream=row_context["stream"] or carried_stream,
                    class_name=row_context["class_name"] or carried_class,
                    batch=row_context["batch"] or carried_batch,
                    subject=subject,
                    topic=topic,
                    teacher=row_context["teacher"],
                    confidence=max(0.05, min(confidence, 0.99)),
                    warnings=warnings,
                    fingerprint=fingerprint,
                )
            )
        return entries

    @staticmethod
    def _sparse_rows(worksheet: Worksheet) -> list[list[Cell]]:
        """Return only populated/hyperlinked cells grouped by row.

        XLSX files exported by Google Sheets often advertise a large rectangular
        grid. ``Worksheet.iter_rows`` materializes every coordinate in that grid,
        including empty cells. Reading the workbook's populated cell store keeps
        sync time proportional to real data rather than sheet dimensions.
        """

        grouped: dict[int, list[Cell]] = defaultdict(list)
        cell_store = getattr(worksheet, "_cells", {})
        candidates = cell_store.values() if cell_store else (
            cell for row in worksheet.iter_rows() for cell in row
        )
        for cell in candidates:
            if isinstance(cell, MergedCell):
                continue
            text = normalize_text(cell.value)
            hyperlink = cell.hyperlink.target if cell.hyperlink else None
            if not text and not hyperlink:
                continue
            grouped[cell.row].append(cell)
        return [
            sorted(grouped[row_number], key=lambda item: item.column)
            for row_number in sorted(grouped)
        ]

    @staticmethod
    def _dense_values(cells: Sequence[Cell]) -> list[str]:
        if not cells:
            return []
        values = [""] * max(cell.column for cell in cells)
        for cell in cells:
            values[cell.column - 1] = normalize_text(cell.value)
        return values

    def _best_header_map(
        self, worksheet: Worksheet, rows: Sequence[Sequence[Cell]] | None = None
    ) -> dict[int, str]:
        best: tuple[int, dict[int, str]] = (0, {})
        sparse_rows = rows if rows is not None else self._sparse_rows(worksheet)
        for row in sparse_rows:
            if not row or row[0].row > 30:
                continue
            mapping: dict[int, str] = {}
            score = 0
            for cell in row:
                text = normalize_text(cell.value).lower()
                if not text:
                    continue
                mapping[cell.column] = text
                if any(
                    key in text
                    for key in (
                        "date",
                        "time",
                        "class",
                        "batch",
                        "subject",
                        "topic",
                        "chapter",
                        "teacher",
                        "faculty",
                        "video",
                        "pdf",
                        "link",
                    )
                ):
                    score += 2
            if score > best[0]:
                best = (score, mapping)
        return best[1]

    def _context_from_row(
        self, values: list[str], headers: dict[int, str], sheet_title: str
    ) -> dict[str, str | None]:
        combined = " | ".join(value for value in values if value)
        result = self._context_from_text(f"{sheet_title} | {combined}")
        for index, value in enumerate(values, start=1):
            if not value:
                continue
            header = headers.get(index, "")
            lower_header = header.lower()
            if "subject" in lower_header:
                result["subject"] = detect_subject(value) or value
            elif "topic" in lower_header or "chapter" in lower_header or "lecture" in lower_header:
                result["topic"] = value
            elif (
                "teacher" in lower_header
                or "faculty" in lower_header
                or "made by" in lower_header
                or "prepared by" in lower_header
            ):
                result["teacher"] = value
            elif (
                "batch" in lower_header
                or lower_header.strip() in {"course", "cource", "course name", "cource name"}
            ):
                result["batch"] = value
            elif "class" in lower_header or "grade" in lower_header:
                result["class_name"] = detect_class(value) or value
        lower_combined = combined.casefold()
        if "humanities" in lower_combined and "commerce" in lower_combined:
            result["stream"] = "Humanities + Commerce"
        elif "science" in lower_combined and "commerce" in lower_combined:
            result["stream"] = "Science + Commerce"

        if not result["topic"]:
            subject = result["subject"]
            candidates = [
                value
                for value in values
                if len(value) >= 6
                and value != subject
                and not self._looks_like_time(value, value)
                and not self._looks_like_date(value, value)
                and not find_urls(value)
                and not detect_class(value)
            ]
            if candidates:
                result["topic"] = max(candidates, key=len)
        return result

    @staticmethod
    def _context_from_text(text: str) -> dict[str, str | None]:
        return {
            "stream": detect_stream(text),
            "class_name": detect_class(text),
            "batch": detect_batch(text),
            "subject": detect_subject(text),
            "topic": None,
            "teacher": None,
        }

    @staticmethod
    def _looks_like_content_header(values: Sequence[str]) -> bool:
        """Return True only for a real header row, never for file names.

        Earlier logic counted words contained inside values. A normal row such
        as "Chapter 1 ... Lecture 1.mp4" plus a PDF was therefore mistaken for
        a repeated header and every link in that row was discarded. Header
        detection now requires multiple short, exact column labels.
        """
        exact_headers = {
            "date", "day", "days", "book", "book name", "subject",
            "teacher", "teacher name", "faculty", "chapter", "chapter name",
            "topic", "topic name", "lecture", "lecture name", "category",
            "video", "video link", "video drive link", "pdf", "pdf link",
            "link", "status", "uploading status", "uploading status in app",
            "uploaded by", "app uploaded by", "remarks",
        }
        normalized: list[str] = []
        for value in values:
            text = normalize_text(value).casefold()
            text = re.sub(r"[^a-z0-9 ]+", " ", text)
            text = re.sub(r"\s+", " ", text).strip()
            if text:
                normalized.append(text)
        hits = sum(1 for value in normalized if value in exact_headers)
        return hits >= 2

    @staticmethod
    def _cells_have_header_words(cells: Iterable[Cell]) -> bool:
        text = " ".join(normalize_text(cell.value).lower() for cell in cells)
        return any(
            word in text
            for word in ("date", "time", "class", "subject", "topic", "teacher", "video", "pdf")
        )

    @staticmethod
    def _style_key(cell: Cell) -> str:
        return stable_hash(
            cell.style_id,
            cell.number_format,
            cell.fill.fill_type,
            cell.font.bold,
            cell.font.italic,
            cell.alignment.horizontal,
            length=16,
        )

    @staticmethod
    def _looks_like_date(value: object, text: str) -> bool:
        return isinstance(value, (date, datetime)) or bool(DATE_RE.search(text))

    @staticmethod
    def _looks_like_time(value: object, text: str) -> bool:
        return isinstance(value, (time, datetime)) or bool(TIME_RE.search(text) or TIME_24_RE.search(text))

    def _find_date(self, cells: Iterable[Cell]) -> date | None:
        for cell in cells:
            value = cell.value
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            text = normalize_text(value)
            match = DATE_RE.search(text)
            if match:
                for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y", "%d.%m.%y"):
                    try:
                        return datetime.strptime(match.group(0), fmt).date()
                    except ValueError:
                        continue
        return None

    def _find_times(self, cells: Iterable[Cell]) -> list[time]:
        found: list[time] = []
        for cell in cells:
            value = cell.value
            if isinstance(value, datetime):
                candidate = value.time().replace(second=0, microsecond=0)
                # Excel date cells are commonly loaded as midnight datetimes.
                # Do not mistake the date's 00:00 component for a class time.
                if candidate != time(0, 0):
                    found.append(candidate)
                continue
            if isinstance(value, time):
                found.append(value.replace(second=0, microsecond=0))
                continue
            text = normalize_text(value)
            for match in TIME_RE.finditer(text):
                hour = int(match.group("hour"))
                minute = int(match.group("minute") or 0)
                meridiem = match.group("meridiem").lower().replace(".", "")
                if meridiem == "pm" and hour != 12:
                    hour += 12
                if meridiem == "am" and hour == 12:
                    hour = 0
                found.append(time(hour, minute))
            for match in TIME_24_RE.finditer(text):
                candidate = time(int(match.group("hour")), int(match.group("minute")))
                if candidate not in found:
                    found.append(candidate)
        return found[:2]

    @staticmethod
    def _nearest_title(worksheet: Worksheet, cell: Cell) -> str:
        cell_store = getattr(worksheet, "_cells", {})
        for column in range(max(1, cell.column - 3), cell.column + 1):
            candidate = cell_store.get((cell.row, column))
            value = normalize_text(candidate.value) if candidate else ""
            if value and not find_urls(value):
                return value
        return ""

    @staticmethod
    def _dedupe_timetable(entries: list[TimetableEntry]) -> list[TimetableEntry]:
        best: dict[str, TimetableEntry] = {}
        for entry in entries:
            existing = best.get(entry.fingerprint)
            if existing is None or entry.confidence > existing.confidence:
                best[entry.fingerprint] = entry
        return sorted(
            best.values(),
            key=lambda item: (
                item.class_date or date.max,
                item.start_time or time.max,
                item.stream or "",
                item.class_name or "",
            ),
        )

    @staticmethod
    def _dedupe_content(items: list[ContentItem]) -> list[ContentItem]:
        best: dict[str, ContentItem] = {}
        for item in items:
            existing = best.get(item.fingerprint)
            if existing is None or item.confidence > existing.confidence:
                best[item.fingerprint] = item
        return sorted(best.values(), key=lambda item: (item.sheet_name, item.source_row, item.title))


def _normalize_url(url: str) -> str:
    return normalize_text(url).rstrip("/").casefold()


def _clean_content_title(title: str) -> str | None:
    value = normalize_text(title)
    if not value or value.startswith("http"):
        return None
    value = re.sub(r"\.(?:mp4|pdf|mov|webm|mkv|docx?|pptx?|xlsx?)\b", "", value, flags=re.I)
    return normalize_text(value) or None


def _subject_from_teacher(teacher: str | None, stream: str | None) -> str | None:
    if stream != "Nirmaan" or not teacher:
        return None
    normalized = normalize_text(teacher).casefold()
    mapping = {
        "kishore": "Mathematics",
        "kishor": "Mathematics",
        "pranav": "Social Science",
        "magnolia": "English",
        "magnoliya": "English",
        "lali": "Hindi",
        "laali": "Hindi",
        "pankaj": "Sanskrit",
    }
    return next((subject for token, subject in mapping.items() if token in normalized), None)


def _subject_from_tab_topic(sheet_title: str, title: str, topic: str | None) -> str | None:
    text = normalize_text(f"{sheet_title} {title} {topic or ''}").casefold()
    tab_rules = {
        "class 7th kishore": "Mathematics",
        "class 7th kishor": "Mathematics",
        "class 7th pranav": "Social Science",
        "class 7th anjlai": "Science",
        "class 7th anjali": "Science",
    }
    for token, subject in tab_rules.items():
        if token in text:
            return subject
    phrase_rules = {
        "electric charge": "Physics",
        "electrostatic": "Physics",
        "current electricity": "Physics",
        "acidic, basic and neutral": "Science",
        "exploring substances": "Science",
        "large numbers around us": "Mathematics",
        "arithmetic expressions": "Mathematics",
        "geographical diversity": "Social Science",
        "night vision": "Science",
        "honey never spoils": "Science",
        "space travel": "Science",
        "osmosis": "Science",
        "gravity": "Science",
        "gavrity": "Science",
        "holes in a bottle": "Science",
        "particulate nature of matter": "Science",
        "study methods": "General",
        "diagram in a": "General",
    }
    for token, subject in phrase_rules.items():
        if token in text:
            return subject
    if "orientation" in text and "nirmaan" in text:
        return "General"
    return None


def _skip_non_learning_link(header: str, url: str) -> bool:
    normalized = normalize_text(header).casefold()
    exact_non_content = {
        "otp", "remarks", "remark", "uploaded by", "app uploaded by",
        "uploading status", "uploading status in app", "pdf uploading status in app",
        "video id", "id", "status", "dpp status", "pdf status",
    }
    if normalized in exact_non_content:
        return True
    lower_url = url.casefold()
    if "admin.educrypt.ai/auth_panel" in lower_url:
        return True
    if "test-series.nexttoppers.in" in lower_url and any(
        token in lower_url for token in ("join", "login", "signup", "register")
    ):
        return True
    return False

def classify_content_type(url: str, label: str = "") -> ContentType:
    lower_url = url.lower()
    lower_label = normalize_text(label).lower()
    lower = f"{lower_url} {lower_label}"
    if "youtube.com" in lower or "youtu.be" in lower:
        return ContentType.VIDEO
    if any(extension in lower_label for extension in (".mp4", ".mov", ".webm", ".mkv")):
        return ContentType.VIDEO
    if any(word in lower_label for word in ("video", "vidoe", "vedio", "recording")):
        return ContentType.VIDEO
    if ".pdf" in lower_label or re.search(r"\bpdf\b", lower_label):
        return ContentType.PDF
    if "canva.com/design" in lower or "canva.link" in lower or re.search(r"\bpptx?\b", lower_label):
        return ContentType.PRESENTATION
    if "wpp" in lower_label or "dpp" in lower_label:
        return ContentType.DOCUMENT
    if "docs.google.com/forms" in lower:
        return ContentType.FORM
    if "docs.google.com/spreadsheets" in lower:
        return ContentType.SPREADSHEET
    if "docs.google.com/presentation" in lower:
        return ContentType.PRESENTATION
    if "docs.google.com/document" in lower:
        return ContentType.DOCUMENT
    url_path = lower_url.split("?", 1)[0]
    for extension, content_type in CONTENT_EXTENSIONS.items():
        if extension in url_path or extension in lower_label:
            return content_type
    if "pdf" in lower:
        return ContentType.PDF
    if any(word in lower for word in ("video", "vidoe", "vedio", "lecture", "recording")):
        return ContentType.VIDEO
    return ContentType.EXTERNAL_LINK

def detect_subject(text: str) -> str | None:
    lower = normalize_text(text).lower()
    matches: list[tuple[int, str]] = []
    for canonical, aliases in SUBJECT_ALIASES.items():
        for alias in aliases:
            if re.search(rf"(?<!\w){re.escape(alias)}(?!\w)", lower):
                matches.append((len(alias), canonical))
    return max(matches, default=(0, None))[1]


def detect_stream(text: str) -> str | None:
    lower = normalize_text(text).lower()
    scores: Counter[str] = Counter()
    for canonical, aliases in STREAM_ALIASES.items():
        for alias in aliases:
            if alias in lower:
                scores[canonical] += len(alias)
    if not scores:
        return None
    return scores.most_common(1)[0][0]


def detect_class(text: str) -> str | None:
    normalized = normalize_text(text)
    explicit = re.search(r"\b(?:class|grade|std\.?|standard)[\s-]*(7|8|9|10|11|12)(?:th|st|nd|rd)?\b", normalized, re.I)
    if explicit:
        return f"Class {explicit.group(1)}"
    # Headings also use forms such as "Prarambh Pro 11th Science". The
    # following match requires a stream word, so lecture numbers inside topic
    # text are not treated as classes.
    stream_heading = re.search(
        r"\b(7|8|9|10|11|12)(?:th|st|nd|rd)?\s+(?:science|commerce|humanities|arts)\b",
        normalized,
        re.I,
    )
    if stream_heading:
        return f"Class {stream_heading.group(1)}"
    subject_tab = re.search(
        r"^(?:hindi|english|maths?|mathematics|physics|chemistry|biology|history|geography|psychology|sociology|accountancy|accounts?|economics|business studies|political science)\s*[-–]?\s*(7|8|9|10|11|12)(?:th|st|nd|rd)?\s*$",
        normalized,
        re.I,
    )
    if subject_tab:
        return f"Class {subject_tab.group(1)}"
    # Tab names commonly use forms such as "Physics - 11".
    suffix = re.search(r"(?:^|[-–|])\s*(7|8|9|10|11|12)(?:th|st|nd|rd)?\s*$", normalized, re.I)
    return f"Class {suffix.group(1)}" if suffix else None


def detect_batch(text: str) -> str | None:
    normalized = normalize_text(text)
    match = re.search(r"\b([A-Za-z][A-Za-z0-9 ._-]{1,30}(?:batch|pro|plus|2\.0))\b", normalized, re.I)
    return match.group(1).strip() if match else None
