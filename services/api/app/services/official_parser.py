from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from typing import Iterable

from openpyxl import Workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.worksheet import Worksheet

from ..models import TimetableEntry
from ..timezones import INDIA_TIMEZONE
from ..utils import normalize_text, stable_hash

OFFICIAL_SOURCES = {
    "1vEGv2nzGYlFWtHT8SmvzGIYeL_Yl9RJNRohM2VHY0Lg": "commerce",
    "1j1yR4DRzqJwnkkfHmeWuv7vP-iDP0KZVahV4l9lcjEg": "science",
    "1M8mxrLzxZ9AixQ7fejbzIjhZGLgNzS7flMx6cXgiXHQ": "humanities",
    "1a7qJpddbGL3fnR6RBmsrFxqjjVLB_g9Ryz7cQ1i0W8U": "nirmaan",
}

WEEKDAYS = {day.lower(): day for day in ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")}
EMPTY_MARKERS = {"", "-", "--", "holiday", "leave", "recorded lectures cuet", "cuet recorded lecture"}
MONTHS = {
    "january": 1, "jan": 1,
    "february": 2, "feb": 2,
    "march": 3, "mar": 3,
    "april": 4, "apr": 4,
    "may": 5,
    "june": 6, "jun": 6,
    "july": 7, "jul": 7,
    "august": 8, "aug": 8,
    "september": 9, "sep": 9, "sept": 9,
    "october": 10, "oct": 10,
    "november": 11, "nov": 11,
    "december": 12, "dec": 12,
}
TIME_RANGE_RE = re.compile(
    r"(?P<h1>\d{1,2})(?::(?P<m1>\d{2}))?\s*(?P<a1>am|pm)?\s*[-–]\s*"
    r"(?P<h2>\d{1,2})(?::(?P<m2>\d{2}))?\s*(?P<a2>am|pm)?",
    re.I,
)

SUBJECT_PATTERNS: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"\b(?:accountancy|accounts?)\b", re.I), "Accountancy"),
    (re.compile(r"\b(?:business studies|business stuides|bst)\b", re.I), "Business Studies"),
    (re.compile(r"\bentrepreneurship\b", re.I), "Entrepreneurship"),
    (re.compile(r"\b(?:applied maths?|applied mathematics)\b", re.I), "Applied Mathematics"),
    (re.compile(r"\b(?:core maths?|mathematics|maths?)\b", re.I), "Mathematics"),
    (re.compile(r"\beconomics?\b", re.I), "Economics"),
    (re.compile(r"\bphysics\b", re.I), "Physics"),
    (re.compile(r"\bchemistry\b", re.I), "Chemistry"),
    (re.compile(r"\bbiology\b", re.I), "Biology"),
    (re.compile(r"\b(?:political science|pol\.?\s*science|political)\b", re.I), "Political Science"),
    (re.compile(r"\bhistory\b", re.I), "History"),
    (re.compile(r"\b(?:geography|geoegraphy)\b|\bgeo\b", re.I), "Geography"),
    (re.compile(r"\bpsychology\b", re.I), "Psychology"),
    (re.compile(r"\bsociology\b", re.I), "Sociology"),
    (re.compile(r"\benglish\b", re.I), "English"),
    (re.compile(r"\bhindi\b", re.I), "Hindi"),
    (re.compile(r"\bsanskrit\b", re.I), "Sanskrit"),
    (re.compile(r"\b(?:social science|sst)\b", re.I), "Social Science"),
    (re.compile(r"\bscience\b", re.I), "Science"),
    (re.compile(r"\bphysical education\b", re.I), "Physical Education"),
]


def source_kind(spreadsheet_id: str | None) -> str | None:
    return OFFICIAL_SOURCES.get(spreadsheet_id or "")


def source_stream(kind: str | None) -> str | None:
    return {
        "commerce": "Commerce",
        "science": "Science",
        "humanities": "Humanities",
        "nirmaan": "Nirmaan",
    }.get(kind or "")


def parse_official_timetable(
    workbook: Workbook,
    *,
    spreadsheet_id: str,
    source_id: str,
    workbook_name: str,
) -> list[TimetableEntry] | None:
    kind = source_kind(spreadsheet_id)
    if kind == "commerce":
        return _parse_commerce(workbook, source_id, workbook_name)
    if kind == "science":
        return _parse_science(workbook, source_id, workbook_name)
    if kind == "humanities":
        return _parse_humanities(workbook, source_id, workbook_name)
    if kind == "nirmaan":
        return _parse_nirmaan(workbook, source_id, workbook_name)
    return None


def _sheet(workbook: Workbook, names: Iterable[str]) -> Worksheet | None:
    normalized = {normalize_text(name).casefold(): name for name in workbook.sheetnames}
    for name in names:
        exact = normalized.get(normalize_text(name).casefold())
        if exact:
            return workbook[exact]
    for actual in workbook.sheetnames:
        lowered = normalize_text(actual).casefold()
        if any(normalize_text(name).casefold() in lowered for name in names):
            return workbook[actual]
    return None


def _current_week_dates() -> list[date]:
    today = datetime.now(INDIA_TIMEZONE).date()
    monday = today - timedelta(days=today.weekday())
    return [monday + timedelta(days=offset) for offset in range(7)]


def _time_bounds(value: object, *, evening: bool = True) -> tuple[time | None, time | None]:
    start = _time_value(value, evening=evening)
    if start is None:
        return None, None
    text = normalize_text(value).casefold()
    match = TIME_RANGE_RE.search(text)
    if not match:
        return start, _default_end(start)
    groups = match.groupdict()
    hour = int(groups["h2"])
    minute = int(groups["m2"] or 0)
    meridiem = (groups["a2"] or groups["a1"] or "").casefold()
    if meridiem == "pm" and hour < 12:
        hour += 12
    elif meridiem == "am" and hour == 12:
        hour = 0
    elif evening and hour < 12:
        hour += 12
    try:
        return start, time(hour, minute)
    except ValueError:
        return start, _default_end(start)


def _subject_list(raw: str) -> list[str]:
    group = canonical_subject_group(raw)
    subjects = [normalize_text(item) for item in group.split(" / ") if normalize_text(item)]
    return subjects or ([normalize_text(raw)] if normalize_text(raw) else [])


def _split_parallel_topics(topic: str) -> list[tuple[str | None, str]]:
    """Split option/parallel lessons without breaking normal multi-line titles."""
    pieces = [normalize_text(item) for item in re.split(r"\s*/\s*|\n+", topic) if normalize_text(item)]
    pieces = [normalize_text(item) for item in re.split(r"\s*/\s*|\n+", topic) if normalize_text(item)]
    if len(pieces) < 2:
        return [(subjects_from_topic(topic), normalize_text(topic))]
    lessons: list[tuple[str | None, str]] = []
    carried: list[str] = []
    for piece in pieces:
        carried.append(piece)
        combined = normalize_text(" ".join(carried))
        subject = subjects_from_topic(combined)
        if subject:
            lessons.append((subject, combined))
            carried = []
    if carried:
        suffix = normalize_text(" ".join(carried))
        if lessons:
            subject, previous = lessons[-1]
            lessons[-1] = (subject, normalize_text(f"{previous} {suffix}"))
        else:
            return [(subjects_from_topic(topic), normalize_text(topic))]
    unique_subjects = {subject for subject, _ in lessons if subject}
    if len(lessons) >= 2 and len(unique_subjects) >= 2:
        return lessons
    return [(subjects_from_topic(topic), normalize_text(topic))]


def _append_recurring_fallback(
    entries: list[TimetableEntry],
    covered: set[tuple[date, str, str]],
    schedules: list[dict[str, object]],
    *,
    source_id: str,
    workbook_name: str,
    stream: str,
) -> None:
    dated = [item.class_date for item in entries if item.class_date]
    if not dated:
        return
    today = datetime.now(INDIA_TIMEZONE).date()
    latest = max(dated)
    # Only bridge a recent planner gap. Old archived workbooks must not generate
    # classes for the current week.
    if abs((today - latest).days) > 21:
        return
    for class_date in _current_week_dates():
        weekday = class_date.strftime("%A")
        for item in schedules:
            class_name = str(item["class_name"])
            batch = str(item["batch"])
            if str(item["day"]) != weekday or (class_date, class_name, batch) in covered:
                continue
            start = item["start"]
            end = item["end"]
            subject = str(item["subject"])
            if not isinstance(start, time) or not isinstance(end, time):
                continue
            entries.append(
                _entry(
                    source_id=source_id,
                    workbook_name=workbook_name,
                    sheet_name=str(item["sheet_name"]),
                    source_row=int(item["source_row"]),
                    coordinates=list(item.get("coordinates", [])),
                    class_date=class_date,
                    day=weekday,
                    start=start,
                    end=end,
                    stream=stream,
                    class_name=class_name,
                    batch=batch,
                    subject=subject,
                    topic="Weekly Planner not updated — topic pending",
                    confidence=0.91,
                    warnings=[
                        "The dated Weekly Planner has no row for this class. Subject and time came from the Time Table tab."
                    ],
                )
            )


def _commerce_recurring_schedule(ws: Worksheet | None) -> list[dict[str, object]]:
    if ws is None:
        return []
    schedules: list[dict[str, object]] = []
    for heading_row in range(1, min(ws.max_row, 60) + 1):
        heading_cols = [
            col for col in range(1, ws.max_column + 1)
            if "class" in normalize_text(ws.cell(heading_row, col).value).casefold()
            and "commerce" in normalize_text(ws.cell(heading_row, col).value).casefold()
        ]
        if not heading_cols:
            continue
        for index, heading_col in enumerate(heading_cols):
            heading = normalize_text(ws.cell(heading_row, heading_col).value)
            class_name = _class_from_text(heading)
            if not class_name:
                continue
            header_row = heading_row + 1
            if normalize_text(ws.cell(header_row, heading_col).value).casefold() not in {"day", "days"}:
                continue
            region_end = heading_cols[index + 1] - 1 if index + 1 < len(heading_cols) else min(ws.max_column, heading_col + 5)
            time_columns: list[tuple[int, time, time]] = []
            for col in range(heading_col + 1, region_end + 1):
                start, end = _time_bounds(ws.cell(header_row, col).value, evening=True)
                if start and end:
                    time_columns.append((col, start, end))
            for row in range(header_row + 1, min(ws.max_row, header_row + 8) + 1):
                day = _weekday(ws.cell(row, heading_col).value)
                if not day:
                    continue
                for col, start, end in time_columns:
                    raw = normalize_text(ws.cell(row, col).value)
                    if _skip_class(raw):
                        continue
                    for subject in _subject_list(raw):
                        schedules.append({
                            "class_name": class_name,
                            "batch": "Prarambh",
                            "day": day,
                            "start": start,
                            "end": end,
                            "subject": subject,
                            "sheet_name": ws.title,
                            "source_row": row,
                            "coordinates": [ws.cell(row, col).coordinate],
                        })
        if schedules:
            break
    return schedules


def _commerce_heading_for_block(ws: Worksheet, header_row: int, start_col: int) -> str:
    """Find the class heading nearest to one Commerce date/day block."""
    for candidate_row in range(header_row - 1, max(0, header_row - 6), -1):
        # Prefer cells around this block so side-by-side Class 11/Class 12 sections
        # do not accidentally borrow the heading from another block.
        for col in range(max(1, start_col - 1), min(ws.max_column, start_col + 5) + 1):
            value = normalize_text(ws.cell(candidate_row, col).value)
            lower = value.casefold()
            if "commerce" in lower and ("class 11" in lower or "class 12" in lower):
                return value
    return ""


def _commerce_matrix_time_columns(
    ws: Worksheet,
    header_row: int,
    start_col: int,
    end_col: int,
) -> dict[tuple[time, time], int]:
    columns: dict[tuple[time, time], int] = {}
    for col in range(start_col + 2, end_col + 1):
        start, end = _time_bounds(ws.cell(header_row, col).value, evening=True)
        if start and end:
            columns[(start, end)] = col
    return columns


def _commerce_override_slot_time(raw: str, start: time, end: time) -> tuple[time, time]:
    """Honor subject-cell overrides such as 'Applied Maths (4:00 PM Onwards)'."""
    text = normalize_text(raw)
    if not re.search(r"\b\d{1,2}(?::\d{2})?\s*(?:am|pm)\b", text, re.I):
        return start, end
    explicit = _time_value(text, evening=True)
    if explicit is None or explicit == start:
        return start, end
    base_start = datetime.combine(date.today(), start)
    base_end = datetime.combine(date.today(), end)
    if base_end <= base_start:
        base_end += timedelta(days=1)
    duration = base_end - base_start
    explicit_dt = datetime.combine(date.today(), explicit)
    return explicit, (explicit_dt + duration).time().replace(second=0, microsecond=0)


def _commerce_slot_lessons(subject_raw: str, topic_raw: object) -> list[tuple[str, str | None]]:
    subjects = _subject_list(subject_raw)
    if not subjects:
        return []
    raw_topic = "" if topic_raw is None else str(topic_raw)
    topic = normalize_text(raw_topic)
    if not topic:
        return [(subject, None) for subject in subjects]

    # A single subject owns the whole topic text, even when the topic happens
    # to contain a subject word in the chapter title.
    if len(subjects) == 1:
        return [(subjects[0], topic)]

    # Parallel slots such as "BST / Core Maths" may have two topic lines, or
    # only one of the two topic lines. Match explicit topic subject labels first.
    mapping: dict[str, str] = {}
    lessons = _split_parallel_topics(raw_topic)
    for lesson_subject, lesson_topic in lessons:
        if not lesson_subject:
            continue
        for candidate in _subject_list(lesson_subject):
            if candidate in subjects and candidate not in mapping:
                mapping[candidate] = normalize_text(lesson_topic)

    if mapping:
        return [(subject, mapping.get(subject)) for subject in subjects]

    # No explicit label: the text is ambiguous. Preserve the class rather than
    # dropping it; attach the text to the first option and leave the other option
    # visible with an unknown topic.
    return [
        (subject, topic if index == 0 else None)
        for index, subject in enumerate(subjects)
    ]


def _parse_commerce_matrix(
    ws: Worksheet,
    *,
    source_id: str,
    workbook_name: str,
    academic_year: int,
) -> tuple[list[TimetableEntry], set[tuple[date, str, str]]]:
    entries: list[TimetableEntry] = []
    covered: set[tuple[date, str, str]] = set()

    for header_row in range(1, min(ws.max_row, 2000) + 1):
        block_starts = [
            col
            for col in range(1, ws.max_column)
            if normalize_text(ws.cell(header_row, col).value).casefold() == "date"
            and normalize_text(ws.cell(header_row, col + 1).value).casefold() in {"day", "days"}
        ]
        if len(block_starts) < 2:
            continue

        # The current Commerce planner uses mirrored side-by-side blocks:
        # subjects on the left, matching topic/details on the right.
        for index in range(len(block_starts) - 1):
            left = block_starts[index]
            right = block_starts[index + 1]

            left_heading = _commerce_heading_for_block(ws, header_row, left)
            right_heading = _commerce_heading_for_block(ws, header_row, right)
            left_class = _class_from_text(left_heading)
            right_class = _class_from_text(right_heading)
            if left_class not in {"Class 11", "Class 12"} or right_class != left_class:
                continue

            # Avoid treating the next class section as the mirror of this one.
            if abs(right - left) > 12:
                continue

            left_end = right - 1
            next_start = block_starts[index + 2] if index + 2 < len(block_starts) else ws.max_column + 1
            right_end = min(ws.max_column, next_start - 1)

            left_times = _commerce_matrix_time_columns(ws, header_row, left, left_end)
            right_times = _commerce_matrix_time_columns(ws, header_row, right, right_end)
            common_slots = [
                (slot, left_col, right_times[slot])
                for slot, left_col in left_times.items()
                if slot in right_times
            ]
            if not common_slots:
                continue

            batch = "Prarambh Pro" if "pro" in left_heading.casefold() else "Prarambh"

            for row in range(header_row + 1, min(ws.max_row, header_row + 10) + 1):
                left_date = _date_value(ws.cell(row, left).value, academic_year)
                right_date = _date_value(ws.cell(row, right).value, academic_year)
                class_date = left_date or right_date
                day = _weekday(ws.cell(row, left + 1).value) or _weekday(ws.cell(row, right + 1).value)

                if class_date is None and day is None:
                    # A true blank line ends the block; a later class section is
                    # handled by its own header row.
                    row_values = [
                        normalize_text(ws.cell(row, col).value)
                        for col in range(left, min(ws.max_column, right_end) + 1)
                    ]
                    if not any(row_values):
                        break
                    continue

                emitted = 0
                for (slot_start, slot_end), subject_col, topic_col in common_slots:
                    subject_raw = normalize_text(ws.cell(row, subject_col).value)
                    if _skip_class(subject_raw):
                        continue
                    topic_raw = ws.cell(row, topic_col).value
                    start, end = _commerce_override_slot_time(subject_raw, slot_start, slot_end)
                    for subject, topic in _commerce_slot_lessons(subject_raw, topic_raw):
                        entries.append(
                            _entry(
                                source_id=source_id,
                                workbook_name=workbook_name,
                                sheet_name=ws.title,
                                source_row=row,
                                coordinates=[
                                    ws.cell(row, subject_col).coordinate,
                                    ws.cell(row, topic_col).coordinate,
                                ],
                                class_date=class_date,
                                day=day,
                                start=start,
                                end=end,
                                stream="Commerce",
                                class_name=left_class,
                                batch=batch,
                                subject=subject,
                                topic=topic,
                                confidence=0.99 if topic else 0.96,
                                warnings=[] if topic else ["Topic not provided in the Commerce planner detail block."],
                            )
                        )
                        emitted += 1

                # Only suppress recurring fallback when the dated planner row
                # actually produced at least one class.
                if class_date is not None and emitted:
                    covered.add((class_date, left_class, batch))

            # Consume pairs as [subject block, detail block], not overlapping
            # [detail block, next block].
            break

    return entries, covered


def _parse_commerce(workbook: Workbook, source_id: str, workbook_name: str) -> list[TimetableEntry]:
    ws = _sheet(workbook, ("Weekly Planner",))
    timetable = _sheet(workbook, ("Time Table/Planner", "Time TablePlanner"))
    if ws is None:
        return []

    academic_year = _academic_start_year(workbook_name)
    entries, covered = _parse_commerce_matrix(
        ws,
        source_id=source_id,
        workbook_name=workbook_name,
        academic_year=academic_year,
    )

    # Keep support for the previous Commerce Weekly Planner layout. This makes
    # the fix additive rather than replacing known working behavior.
    for header_row in range(1, min(ws.max_row, 2000) + 1):
        if normalize_text(ws.cell(header_row, 1).value).casefold() != "date":
            continue
        if normalize_text(ws.cell(header_row, 2).value).casefold() not in {"day", "days"}:
            continue
        class_text = _nearest_heading(ws, header_row, ("class 11", "class 12"))
        class_name = _class_from_text(class_text)
        if class_name not in {"Class 11", "Class 12"}:
            continue
        batch = "Prarambh Pro" if "pro" in class_text.casefold() else "Prarambh"
        pairs: list[tuple[int, int, time, time]] = []
        for col in range(3, ws.max_column):
            if normalize_text(ws.cell(header_row, col).value).casefold() != "subject":
                continue
            start, end = _time_bounds(ws.cell(header_row, col + 1).value, evening=True)
            if start and end:
                pairs.append((col, col + 1, start, end))
        if not pairs:
            continue
        for row in range(header_row + 1, min(ws.max_row, header_row + 10) + 1):
            class_date = _date_value(ws.cell(row, 1).value, academic_year)
            day = _weekday(ws.cell(row, 2).value)
            if class_date is None and day is None:
                break
            # When the newer side-by-side matrix has already emitted this
            # date/class, it is the richer authoritative record. Do not also
            # ingest the older subject-only block for the same class/date.
            if class_date is not None and (class_date, class_name, batch) in covered:
                continue
            emitted = 0
            for subject_col, topic_col, start, end in pairs:
                subject_raw = normalize_text(ws.cell(row, subject_col).value)
                topic = normalize_text(ws.cell(row, topic_col).value)
                if _skip_class(subject_raw):
                    continue
                if not topic:
                    for subject in _subject_list(subject_raw):
                        entries.append(
                            _entry(
                                source_id=source_id,
                                workbook_name=workbook_name,
                                sheet_name=ws.title,
                                source_row=row,
                                coordinates=[ws.cell(row, subject_col).coordinate, ws.cell(row, topic_col).coordinate],
                                class_date=class_date,
                                day=day,
                                start=start,
                                end=end,
                                stream="Commerce",
                                class_name=class_name,
                                batch=batch,
                                subject=subject,
                                topic=None,
                                confidence=0.96,
                                warnings=["Topic not provided in the Commerce planner detail block."],
                            )
                        )
                        emitted += 1
                    continue
                lessons = _split_parallel_topics(topic)
                if len(lessons) == 1 and lessons[0][0] is None:
                    subjects = _subject_list(subject_raw)
                    lessons = [(" / ".join(subjects), topic)]
                for lesson_subject, lesson_topic in lessons:
                    subject = lesson_subject or canonical_subject_group(subject_raw)
                    entries.append(_entry(
                        source_id=source_id, workbook_name=workbook_name, sheet_name=ws.title,
                        source_row=row, coordinates=[ws.cell(row, subject_col).coordinate, ws.cell(row, topic_col).coordinate],
                        class_date=class_date, day=day, start=start, end=end, stream="Commerce",
                        class_name=class_name, batch=batch, subject=subject, topic=lesson_topic,
                        confidence=0.99, warnings=[],
                    ))
                    emitted += 1
            if class_date is not None and emitted:
                covered.add((class_date, class_name, batch))

    _append_recurring_fallback(
        entries,
        covered,
        _commerce_recurring_schedule(timetable),
        source_id=source_id,
        workbook_name=workbook_name,
        stream="Commerce",
    )
    return _merge_duplicate_batches(entries)

def _parse_science(workbook: Workbook, source_id: str, workbook_name: str) -> list[TimetableEntry]:
    planner = _sheet(workbook, ("Weekly Planner",))
    timetable = _sheet(workbook, ("Time Table", "Timetable"))
    if planner is None:
        return []
    schedule = _science_schedule(timetable) if timetable else {}
    topic_index = _subject_topic_index(
        workbook,
        allowed={"Physics", "Chemistry", "Biology", "Mathematics"},
    )
    academic_year = _academic_start_year(workbook_name)
    blocks: list[tuple[int, int, str, str]] = []
    for col in range(1, planner.max_column + 1):
        heading = normalize_text(planner.cell(1, col).value)
        if "science" not in heading.casefold() or "class" not in heading.casefold():
            continue
        class_name = _class_from_text(heading)
        if not class_name:
            continue
        batch = "Prarambh Pro" if "pro" in heading.casefold() else "Prarambh"
        next_heading = next((c for c in range(col + 1, planner.max_column + 1)
                            if "class" in normalize_text(planner.cell(1, c).value).casefold()), planner.max_column + 1)
        blocks.append((col, next_heading - 1, class_name, batch))
    entries: list[TimetableEntry] = []
    covered: set[tuple[date, str, str]] = set()
    for start_col, end_col, class_name, batch in blocks:
        date_col, day_col = start_col, start_col + 1
        topic_cols = [c for c in range(start_col + 2, end_col + 1) if _time_value(planner.cell(2, c).value, evening=True)]
        for row in range(3, planner.max_row + 1):
            class_date = _date_value(planner.cell(row, date_col).value, academic_year)
            day = _weekday(planner.cell(row, day_col).value)
            if class_date is None and day is None:
                continue
            if class_date is not None:
                covered.add((class_date, class_name, batch))
            for col in topic_cols:
                topic = normalize_text(planner.cell(row, col).value)
                if _skip_class(topic):
                    continue
                start, end = _time_bounds(planner.cell(2, col).value, evening=True)
                if start is None or end is None:
                    continue
                lessons = _split_parallel_topics(topic)
                for detected_subject, lesson_topic in lessons:
                    subject = (
                        detected_subject
                        or _subjects_from_topic_index(lesson_topic, topic_index)
                        or schedule.get((class_name, batch, day or "", start))
                    )
                    warnings = [] if subject else ["Subject could not be matched to the timetable pattern"]
                    entries.append(_entry(
                        source_id=source_id, workbook_name=workbook_name, sheet_name=planner.title,
                        source_row=row, coordinates=[planner.cell(row, date_col).coordinate, planner.cell(row, col).coordinate],
                        class_date=class_date, day=day, start=start, end=end, stream="Science",
                        class_name=class_name, batch=batch, subject=subject, topic=lesson_topic,
                        confidence=0.99 if subject else 0.76, warnings=warnings,
                    ))
    recurring: list[dict[str, object]] = []
    for (class_name, batch, day, start), subject_group in schedule.items():
        for subject in _subject_list(subject_group):
            recurring.append({"class_name": class_name, "batch": batch, "day": day, "start": start,
                              "end": _default_end(start), "subject": subject, "sheet_name": timetable.title if timetable else "Time Table",
                              "source_row": 1, "coordinates": []})
    _append_recurring_fallback(entries, covered, recurring, source_id=source_id, workbook_name=workbook_name, stream="Science")
    return _merge_duplicate_batches(entries)

def _science_schedule(ws: Worksheet) -> dict[tuple[str, str, str, time], str]:
    result: dict[tuple[str, str, str, time], str] = {}
    heading_cols = [
        col
        for col in range(1, min(ws.max_column, 30) + 1)
        if "science" in normalize_text(ws.cell(1, col).value).casefold()
        and "class" in normalize_text(ws.cell(1, col).value).casefold()
    ]
    for index, heading_col in enumerate(heading_cols):
        heading = normalize_text(ws.cell(1, heading_col).value)
        class_name = _class_from_text(heading)
        if not class_name:
            continue
        batch = "Prarambh Pro" if "pro" in heading.casefold() else "Prarambh"
        if normalize_text(ws.cell(2, heading_col).value).casefold() not in {"day", "days"}:
            continue
        region_end = heading_cols[index + 1] - 1 if index + 1 < len(heading_cols) else min(ws.max_column, heading_col + 3)
        time_cols: list[tuple[int, time]] = []
        for col in range(heading_col + 1, region_end + 1):
            start = _time_value(ws.cell(2, col).value, evening=True)
            if start:
                time_cols.append((col, start))
        for row in range(3, min(ws.max_row, 9) + 1):
            day = _weekday(ws.cell(row, heading_col).value)
            if not day:
                continue
            for col, start in time_cols:
                subject_raw = normalize_text(ws.cell(row, col).value)
                if _skip_class(subject_raw):
                    continue
                result[(class_name, batch, day, start)] = canonical_subject_group(subject_raw)
    return result


def _parse_humanities(workbook: Workbook, source_id: str, workbook_name: str) -> list[TimetableEntry]:
    planner = _sheet(workbook, ("Weekly Planner", "Weekly "))
    timetable = _sheet(workbook, ("Time Table master", "Time Table Master"))
    if planner is None:
        return []
    schedule = _humanities_schedule(timetable) if timetable else {}
    academic_year = _academic_start_year(workbook_name)
    headings: list[tuple[int, int, str, str]] = []
    heading_cols = [c for c in range(1, planner.max_column + 1)
                    if "weekly planner" in normalize_text(planner.cell(1, c).value).casefold()]
    for index, col in enumerate(heading_cols):
        heading = normalize_text(planner.cell(1, col).value)
        class_name = _class_from_text(heading)
        if not class_name:
            continue
        batch = "Prarambh Pro" if "pro" in heading.casefold() else "Prarambh"
        next_col = heading_cols[index + 1] if index + 1 < len(heading_cols) else planner.max_column + 1
        headings.append((col, next_col - 1, class_name, batch))
    entries: list[TimetableEntry] = []
    covered: set[tuple[date, str, str]] = set()
    for start_col, end_col, class_name, batch in headings:
        topic_cols = [c for c in range(start_col, end_col + 1) if _time_value(planner.cell(2, c).value, evening=True)]
        for row in range(3, planner.max_row + 1):
            class_date = _date_value(planner.cell(row, 1).value, academic_year)
            day = _weekday(planner.cell(row, 2).value)
            if class_date is None and day is None:
                continue
            if class_date is not None:
                covered.add((class_date, class_name, batch))
            for col in topic_cols:
                topic = normalize_text(planner.cell(row, col).value)
                if _skip_class(topic):
                    continue
                start, end = _time_bounds(planner.cell(2, col).value, evening=True)
                if start is None or end is None:
                    continue
                if "5:15" in topic.casefold():
                    start = time(17, 15)
                    end = _default_end(start)
                lessons = _split_parallel_topics(topic)
                for detected_subject, lesson_topic in lessons:
                    subject = detected_subject or schedule.get((class_name, batch, day or "", start))
                    warnings = [] if subject else ["Subject could not be matched to the timetable pattern"]
                    entries.append(_entry(
                        source_id=source_id, workbook_name=workbook_name, sheet_name=planner.title,
                        source_row=row, coordinates=[planner.cell(row, 1).coordinate, planner.cell(row, col).coordinate],
                        class_date=class_date, day=day, start=start, end=end, stream="Humanities",
                        class_name=class_name, batch=batch, subject=subject, topic=lesson_topic,
                        confidence=0.99 if subject else 0.76, warnings=warnings,
                    ))
    recurring: list[dict[str, object]] = []
    for (class_name, batch, day, start), subject_group in schedule.items():
        for subject in _subject_list(subject_group):
            recurring.append({"class_name": class_name, "batch": batch, "day": day, "start": start,
                              "end": _default_end(start), "subject": subject, "sheet_name": timetable.title if timetable else "Time Table master",
                              "source_row": 1, "coordinates": []})
    _append_recurring_fallback(entries, covered, recurring, source_id=source_id, workbook_name=workbook_name, stream="Humanities")
    return _merge_duplicate_batches(entries)

def _humanities_schedule(ws: Worksheet) -> dict[tuple[str, str, str, time], str]:
    result: dict[tuple[str, str, str, time], str] = {}
    # The sheet contains independent horizontal regions. Each region begins with a Day column.
    for row1 in range(1, min(ws.max_row, 30) + 1):
        day_cols = [c for c in range(1, ws.max_column + 1)
                    if normalize_text(ws.cell(row1, c).value).casefold().startswith("day")]
        if not day_cols:
            continue
        for day_col in day_cols:
            region_end = next((c - 1 for c in day_cols if c > day_col), ws.max_column)
            heading_row = row1 + 1
            time_row = row1 + 2
            class_cols = [c for c in range(day_col + 1, region_end + 1)
                          if "class" in normalize_text(ws.cell(heading_row, c).value).casefold()]
            for idx, class_col in enumerate(class_cols):
                heading = normalize_text(ws.cell(heading_row, class_col).value)
                class_name = _class_from_text(heading)
                if not class_name:
                    continue
                batch = "Prarambh Pro" if "pro" in heading.casefold() else "Prarambh"
                group_end = class_cols[idx + 1] - 1 if idx + 1 < len(class_cols) else region_end
                time_cols = [(c, _time_value(ws.cell(time_row, c).value, evening=True))
                             for c in range(class_col, group_end + 1)]
                time_cols = [(c, t) for c, t in time_cols if t]
                for row in range(time_row + 1, min(ws.max_row, time_row + 8) + 1):
                    day = _weekday(ws.cell(row, day_col).value)
                    if not day:
                        continue
                    for col, start in time_cols:
                        raw = normalize_text(ws.cell(row, col).value)
                        if _skip_class(raw):
                            continue
                        actual_start = time(17, 15) if "5:15" in raw.casefold() else start
                        result[(class_name, batch, day, actual_start)] = canonical_subject_group(raw)
    return result


def _parse_nirmaan(workbook: Workbook, source_id: str, workbook_name: str) -> list[TimetableEntry]:
    """Parse every Nirmaan timetable/planner pair, including NCERT variants.

    The live workbook currently contains independent NEW NCERT, OLD NCERT and
    NEW NCERT 2.0 blocks. Older versions kept only the final pair because a
    dictionary overwrote earlier title rows. This parser keeps each block and
    labels it clearly so Class 8 students never see mixed subjects/topics.
    """
    ws = _sheet(workbook, ("Weekly Planner",))
    if ws is None:
        return []

    academic_year = _academic_start_year(workbook_name)
    events: list[tuple[int, str, str]] = []
    current_batch = "Nirmaan"
    scan_limit = min(ws.max_row, 500)

    for row in range(1, scan_limit + 1):
        row_text = " | ".join(
            normalize_text(ws.cell(row, col).value)
            for col in range(1, ws.max_column + 1)
        )
        lowered = row_text.casefold().replace("2.o", "2.0")
        if "new ncert" in lowered and "2.0" in lowered:
            current_batch = "Nirmaan 2.0"
        elif "old ncert" in lowered:
            current_batch = "Nirmaan Old NCERT"
        elif "new ncert" in lowered:
            current_batch = "Nirmaan New NCERT"

        if "nirmaan" not in lowered:
            continue
        title_batch = "Nirmaan 2.0" if "2.0" in lowered else current_batch
        if "time table" in lowered or "timetable" in lowered:
            events.append((row, "timetable", title_batch))
        elif "planner" in lowered:
            events.append((row, "planner", title_batch))

    entries: list[TimetableEntry] = []
    for index, (title_row, kind, batch) in enumerate(events):
        if kind != "timetable":
            continue
        next_timetable_row = next(
            (row for row, event_kind, _ in events[index + 1 :] if event_kind == "timetable"),
            scan_limit + 1,
        )
        planner_event = next(
            (
                (row, planner_batch)
                for row, event_kind, planner_batch in events[index + 1 :]
                if event_kind == "planner" and row < next_timetable_row
            ),
            None,
        )
        if planner_event is None:
            continue
        planner_row, planner_batch = planner_event
        batch = planner_batch if planner_batch != "Nirmaan" else batch

        timetable_values = _read_nirmaan_block(
            ws, title_row, planner_row, academic_year
        )
        planner_values = _read_nirmaan_block(
            ws, planner_row, next_timetable_row, academic_year
        )

        for key, subject_item in timetable_values.items():
            planner_item = planner_values.get(key)
            if planner_item is None:
                continue
            subject_raw, subject_row, subject_col, day = subject_item
            topic, topic_row, topic_col, planner_day = planner_item
            class_date, start, _lane_column = key
            subject = canonical_subject(subject_raw) or subject_raw
            warnings: list[str] = []
            displayed_day = day or planner_day or class_date.strftime("%A")
            if displayed_day and displayed_day != class_date.strftime("%A"):
                warnings.append("The weekday and date in the Google Sheet do not match")
            entries.append(
                _entry(
                    source_id=source_id,
                    workbook_name=workbook_name,
                    sheet_name=ws.title,
                    source_row=topic_row,
                    coordinates=[
                        ws.cell(subject_row, subject_col).coordinate,
                        ws.cell(topic_row, topic_col).coordinate,
                    ],
                    class_date=class_date,
                    day=displayed_day,
                    start=start,
                    end=_default_end(start),
                    stream="Nirmaan",
                    class_name="Class 8",
                    batch=batch,
                    subject=subject,
                    topic=topic,
                    confidence=0.99,
                    warnings=warnings,
                    fingerprint_batch=True,
                )
            )
    return sorted(
        entries,
        key=lambda item: (
            item.class_date or date.max,
            item.start_time or time.max,
            item.batch or "",
            item.subject or "",
        ),
    )


def _read_nirmaan_block(
    ws: Worksheet,
    title_row: int,
    stop_row: int,
    academic_year: int,
) -> dict[tuple[date, time, int], tuple[str, int, int, str]]:
    """Read one Nirmaan block, including visually merged time headings.

    The live Old NCERT block displays one ``18:00`` heading merged across two
    subject columns. Google/Excel stores the second header cell as blank, so a
    parser that reads only explicit time cells silently drops the second class.
    Active data columns now inherit the nearest time heading to their left, and
    the physical column is part of the key so simultaneous classes never
    overwrite one another.
    """
    header_row = next(
        (
            row
            for row in range(title_row + 1, min(stop_row, title_row + 5))
            if any(
                _time_value(ws.cell(row, col).value, evening=True)
                for col in range(1, ws.max_column + 1)
            )
        ),
        None,
    )
    if header_row is None:
        return {}

    data_start = header_row + 1
    data_stop = min(stop_row, ws.max_row + 1)
    active_columns = {
        col
        for col in range(1, ws.max_column + 1)
        if any(
            not _skip_class(normalize_text(ws.cell(row, col).value))
            for row in range(data_start, data_stop)
        )
    }

    time_columns: list[tuple[int, time]] = []
    inherited_time: time | None = None
    for col in range(1, ws.max_column + 1):
        explicit_time = _time_value(ws.cell(header_row, col).value, evening=True)
        if explicit_time is not None:
            inherited_time = explicit_time
        if col in active_columns and inherited_time is not None:
            # Date/day columns occur before the first time heading, therefore
            # they never enter this list. Blank cells after a merged heading do.
            time_columns.append((col, inherited_time))

    values: dict[tuple[date, time, int], tuple[str, int, int, str]] = {}
    blank_rows = 0
    for row in range(data_start, data_stop):
        raw_values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        texts = [normalize_text(value) for value in raw_values]
        if not any(texts):
            blank_rows += 1
            if blank_rows >= 2:
                break
            continue
        blank_rows = 0
        class_date = next(
            (
                parsed
                for value in raw_values
                if (parsed := _date_value(value, academic_year)) is not None
            ),
            None,
        )
        day = next(
            (
                parsed
                for value in raw_values
                if (parsed := _weekday(value)) is not None
            ),
            None,
        )
        if class_date is None:
            continue
        for col, start in time_columns:
            text = normalize_text(ws.cell(row, col).value)
            if _skip_class(text):
                continue
            values[(class_date, start, col)] = (
                text,
                row,
                col,
                day or class_date.strftime("%A"),
            )
    return values

def canonical_subject(text: str) -> str | None:
    normalized = normalize_text(text)
    for pattern, canonical in SUBJECT_PATTERNS:
        if pattern.search(normalized):
            return canonical
    return None


def _clean_subject_group(found: list[str]) -> list[str]:
    if "Political Science" in found and "Science" in found:
        found.remove("Science")
    if "Applied Mathematics" in found and "Mathematics" in found:
        found.remove("Mathematics")
    return found


def canonical_subject_group(text: str) -> str:
    normalized = normalize_text(text)
    found: list[str] = []
    for pattern, canonical in SUBJECT_PATTERNS:
        if pattern.search(normalized) and canonical not in found:
            found.append(canonical)
    found = _clean_subject_group(found)
    if found:
        return " / ".join(found)
    return normalized


def subjects_from_topic(topic: str) -> str | None:
    found: list[str] = []
    for pattern, canonical in SUBJECT_PATTERNS:
        if pattern.search(topic) and canonical not in found:
            found.append(canonical)
    found = _clean_subject_group(found)
    return " / ".join(found) if found else None


def _topic_key(value: object) -> str:
    text = normalize_text(value).casefold()
    text = re.sub(r"https?://\S+", " ", text)
    text = re.sub(r"\.(?:mp4|pdf|mov|webm|mkv)\b", " ", text)
    text = re.sub(r"\((?:youtube|yt|l(?:ecture)?\s*\d+|free class)[^)]*\)", " ", text)
    text = re.sub(r"\b(?:youtube|lecture|class|video|pdf)\b", " ", text)
    text = re.sub(r"\bl\s*\d+\b", " ", text)
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _subject_topic_index(
    workbook: Workbook,
    *,
    allowed: set[str],
) -> dict[str, set[str]]:
    index: dict[str, set[str]] = defaultdict(set)
    for ws in workbook.worksheets:
        subject = canonical_subject(ws.title)
        if subject not in allowed:
            continue
        # Official subject tabs follow names such as "Physics - 11". This
        # avoids indexing timetable and generic session tabs as subject proof.
        if not re.search(r"(?:class\s*)?(?:7|8|9|10|11|12)(?:st|nd|rd|th)?\s*$", normalize_text(ws.title), re.I):
            continue
        for cell in getattr(ws, "_cells", {}).values():
            key = _topic_key(cell.value)
            if 5 <= len(key) <= 180 and not key.isdigit():
                index[key].add(subject)
    return index


def _subjects_from_topic_index(topic: str, index: dict[str, set[str]]) -> str | None:
    segments = [segment for segment in re.split(r"\s*/\s*|\s+\|\s+", normalize_text(topic)) if segment]
    found: list[str] = []
    for segment in segments:
        key = _topic_key(segment)
        if not key:
            continue
        candidates = set(index.get(key, set()))
        if not candidates:
            # Chapter titles sometimes include a lecture suffix or a short
            # prefix. A short key such as "Sets" is accepted only when every
            # word-boundary match points to the same subject.
            for indexed, subjects in index.items():
                descriptive = len(key) >= 8 and (key in indexed or indexed in key)
                token_match = len(key) >= 4 and re.search(rf"\b{re.escape(key)}\b", indexed)
                if descriptive or token_match:
                    candidates.update(subjects)
        if len(candidates) == 1:
            subject = next(iter(candidates))
            if subject not in found:
                found.append(subject)
    return " / ".join(found) if found else None


def _entry(
    *, source_id: str, workbook_name: str, sheet_name: str, source_row: int,
    coordinates: list[str], class_date: date | None, day: str | None, start: time,
    end: time, stream: str, class_name: str, batch: str, subject: str | None,
    topic: str | None, confidence: float, warnings: list[str],
    fingerprint_batch: bool = False,
) -> TimetableEntry:
    fingerprint = stable_hash(
        class_date, start, end, stream, class_name,
        batch if fingerprint_batch else "", subject, topic,
    )
    return TimetableEntry(
        id=stable_hash("timetable", source_id, sheet_name, source_row, ",".join(coordinates), fingerprint),
        source_id=source_id, workbook_name=workbook_name, sheet_name=sheet_name,
        source_row=source_row, source_coordinates=coordinates, class_date=class_date,
        day=day or (class_date.strftime("%A") if class_date else None), start_time=start,
        end_time=end, stream=stream, class_name=class_name, batch=batch, subject=subject,
        topic=topic, teacher=None, confidence=confidence, warnings=warnings, fingerprint=fingerprint,
    )


def _merge_duplicate_batches(entries: list[TimetableEntry]) -> list[TimetableEntry]:
    grouped: dict[str, list[TimetableEntry]] = defaultdict(list)
    for entry in entries:
        grouped[entry.fingerprint].append(entry)
    merged: list[TimetableEntry] = []
    for fingerprint, group in grouped.items():
        best = max(group, key=lambda item: item.confidence)
        batches = sorted({item.batch for item in group if item.batch})
        if len(batches) > 1:
            best = best.model_copy(update={"batch": " + ".join(batches), "id": stable_hash("merged", best.source_id, fingerprint)})
        merged.append(best)
    return sorted(merged, key=lambda item: (item.class_date or date.max, item.start_time or time.max, item.class_name or "", item.subject or ""))


def _date_value(value: object, academic_start_year: int) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 30000 < float(value) < 70000:
        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    text = normalize_text(value)
    if not text:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    match = re.search(r"(?P<day>\d{1,2})(?:st|nd|rd|th)?\s+(?P<month>[A-Za-z]+)(?:\s+(?P<year>\d{4}))?", text)
    if match:
        month = MONTHS.get(match.group("month").casefold())
        if month:
            year = int(match.group("year")) if match.group("year") else academic_start_year + (1 if month <= 3 else 0)
            try:
                return date(year, month, int(match.group("day")))
            except ValueError:
                return None
    return None


def _time_value(value: object, *, evening: bool = False) -> time | None:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        hour, minute = value.hour, value.minute
        if evening and hour < 12:
            hour += 12
        return time(hour, minute)
    if isinstance(value, (int, float)) and 0 <= float(value) < 1:
        seconds = round(float(value) * 86400)
        hour, rem = divmod(seconds, 3600)
        minute = rem // 60
        if evening and hour < 12:
            hour += 12
        return time(hour % 24, minute)
    text = normalize_text(value).casefold()
    if not text:
        return None
    range_match = TIME_RANGE_RE.search(text)
    target = range_match.groupdict() if range_match else None
    if target:
        hour = int(target["h1"]); minute = int(target["m1"] or 0); meridiem = target["a1"] or ""
    else:
        match = re.search(r"(\d{1,2})(?::(\d{2}))?\s*(am|pm)?", text)
        if not match:
            return None
        hour = int(match.group(1)); minute = int(match.group(2) or 0); meridiem = match.group(3) or ""
    if meridiem == "pm" and hour < 12:
        hour += 12
    elif meridiem == "am" and hour == 12:
        hour = 0
    elif evening and hour < 12:
        hour += 12
    try:
        return time(hour, minute)
    except ValueError:
        return None


def _default_end(start: time) -> time:
    minutes = 75 if start == time(18, 30) else 60 if start == time(19, 0) else 90
    dt = datetime.combine(date.today(), start) + timedelta(minutes=minutes)
    return dt.time().replace(second=0, microsecond=0)


def _weekday(value: object) -> str | None:
    return WEEKDAYS.get(normalize_text(value).casefold())


def _class_from_text(text: str) -> str | None:
    match = re.search(r"\bclass\s*(7|8|9|10|11|12)(?:st|nd|rd|th)?\b", normalize_text(text), re.I)
    return f"Class {match.group(1)}" if match else None


def _nearest_heading(ws: Worksheet, row: int, needles: tuple[str, ...]) -> str:
    for candidate_row in range(row - 1, max(0, row - 5), -1):
        values = [normalize_text(ws.cell(candidate_row, col).value) for col in range(1, min(ws.max_column, 15) + 1)]
        for value in values:
            lower = value.casefold()
            if any(needle in lower for needle in needles):
                return value
    return ""


def _academic_start_year(workbook_name: str) -> int:
    match = re.search(r"(20\d{2})\s*[-–/]\s*(?:20)?\d{2}", workbook_name)
    return int(match.group(1)) if match else 2026


def _skip_class(value: str) -> bool:
    normalized = normalize_text(value).casefold().strip()
    return normalized in EMPTY_MARKERS or normalized.startswith("holiday") or normalized.startswith("leave")
