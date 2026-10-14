from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from ..models import ContentItem, ContentType
from ..utils import find_urls, normalize_search_text, normalize_text, stable_hash
from .official_parser import canonical_subject

_HEADER_RE = re.compile(r"[^a-z0-9]+")

_CLASS_RE = re.compile(r"\b(?:class|grade|std\.?|standard)?[\s-]*(7|8|9|10|11|12)(?:th|st|nd|rd)?\b", re.I)

def _detect_class(value: str) -> str | None:
    match = _CLASS_RE.search(value or "")
    return f"Class {match.group(1)}" if match else None

def _detect_stream(value: str) -> str | None:
    text = (value or "").casefold()
    if any(word in text for word in ("commerce", "accountancy", "accounts", "business studies", "bst")):
        return "Commerce"
    if any(word in text for word in ("science", "physics", "chemistry", "biology", "pcmb", "pcm", "pcb")):
        return "Science"
    if any(word in text for word in ("humanities", "political science", "history", "geography", "psychology", "sociology")):
        return "Humanities"
    if any(word in text for word in ("nirmaan", "nirman", "class 8", "8th")):
        return "Nirmaan"
    return None

def _classify_content_type(url: str, label: str) -> ContentType:
    text = f"{url} {label}".casefold()
    if any(token in text for token in ("youtube.com", "youtu.be", ".mp4", ".mov", ".webm", ".mkv")):
        return ContentType.VIDEO
    if ".pdf" in text:
        return ContentType.PDF
    if any(token in text for token in (".ppt", ".pptx", "docs.google.com/presentation", "canva.com")):
        return ContentType.PRESENTATION
    if any(token in text for token in (".xls", ".xlsx", "docs.google.com/spreadsheets")):
        return ContentType.SPREADSHEET
    if any(token in text for token in (".doc", ".docx", "docs.google.com/document")):
        return ContentType.DOCUMENT
    return ContentType.EXTERNAL_LINK

# These are intentionally broad and spelling-tolerant. New Sheets do not have to
# copy one exact template, but understandable headings make automatic indexing
# more reliable and let the UI explain what is missing.
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("date", "dates", "created date", "upload date"),
    "content_type": ("content type", "type", "media type", "file type"),
    "course": ("course", "cource", "course name", "cource name", "batch", "batch name"),
    "stream": ("stream", "branch"),
    "class_name": ("class", "grade", "standard", "std", "class name"),
    "subject": ("subject", "subject name"),
    "chapter_number": ("chapter number", "chapter no", "chapter", "paper no", "test no", "test number"),
    "topic": ("chapter name", "topic", "lesson", "title", "test name", "paper name", "content title", "149"),
    "teacher": ("made by", "prepared by", "teacher", "faculty", "content owner", "owner"),
    "document": ("docs link", "doc link", "document link", "file link", "document", "doc", "link"),
    "pdf": ("pdf", "pdf link", "pdf file"),
    "video": ("video", "video link", "video drive link", "vidoe drive link", "youtube", "youtube link"),
    "presentation": ("presentation", "ppt", "ppt link", "slides", "canva"),
    "spreadsheet": ("spreadsheet", "sheet link", "excel", "xlsx"),
    "app_status": ("app status", "status in website", "website status", "status"),
    "test_date": ("test date", "due date", "scheduled date"),
    "result_dashboard": ("result dashboard", "result deshbord", "result status"),
    "video_solution": ("video solution update", "video solution upodat", "video solution status"),
    "questions": ("no of question", "no of que", "total questions", "questions"),
    "duration": ("test timing", "test timing in minutes", "duration", "minutes"),
    "marks": ("test marks", "marks", "total marks"),
    "remarks": ("remark", "remarks", "note", "notes", "raw files"),
}

IDENTITY_FIELDS = {"topic", "subject", "class_name", "course", "chapter_number", "teacher"}
FILE_FIELDS = {"document", "pdf", "video", "presentation", "spreadsheet"}
TEST_OPERATION_FIELDS = {"app_status", "test_date", "result_dashboard", "video_solution"}

_FILE_PLACEHOLDERS = {"", "-", "--", "done", "on app", "ok", "yes", "no", "pending", "uploaded", "live", "n a", "na"}

def _usable_structured_file(field_name: str, title: str | None, url: str | None) -> bool:
    if url:
        return True
    value = normalize_text(title).casefold()
    if not value or value in _FILE_PLACEHOLDERS:
        return False
    if field_name == "pdf":
        return ".pdf" in value or value.endswith(" pdf") or "pdf file" in value
    return len(value) >= 4



@dataclass(slots=True)
class SheetSchemaProfile:
    header_row: int | None = None
    columns: dict[str, int] = field(default_factory=dict)
    labels: dict[str, str] = field(default_factory=dict)
    searchable_fields: list[str] = field(default_factory=list)
    missing_recommended_fields: list[str] = field(default_factory=list)
    search_ready: bool = False
    record_count: int = 0
    confidence: float = 0.0


def clean_header(value: object) -> str:
    text = normalize_text(value).casefold().replace("&", " and ")
    return _HEADER_RE.sub(" ", text).strip()


def match_field(header: str) -> str | None:
    if not header:
        return None
    for field_name, aliases in FIELD_ALIASES.items():
        if header in aliases:
            return field_name
    # Tolerate common spelling, extra words, and line breaks.
    if "result" in header and ("dashboard" in header or "deshbord" in header):
        return "result_dashboard"
    if "video" in header and "solution" in header:
        return "video_solution"
    if "app" in header and "status" in header:
        return "app_status"
    if "status" in header and "website" in header:
        return "app_status"
    if "chapter" in header and "name" in header:
        return "topic"
    if "chapter" in header and ("number" in header or header.endswith(" no")):
        return "chapter_number"
    if "made by" in header or "prepared by" in header or "teacher" in header or "faculty" in header:
        return "teacher"
    if "question" in header:
        return "questions"
    if "timing" in header or "duration" in header:
        return "duration"
    if "pdf" in header:
        return "pdf"
    if (("video" in header or "youtube" in header or "vidoe" in header or header.startswith("yt "))
            and any(token in header for token in ("link", "url", "file"))):
        return "video"
    if "doc" in header and "link" in header:
        return "document"
    if "file" in header and "link" in header:
        return "document"
    if "sheet" in header and "link" in header:
        return "spreadsheet"
    return None




def _populated_rows(worksheet: Worksheet) -> dict[int, dict[int, Cell]]:
    """Return only cells that contain a value or hyperlink, grouped by row.

    Google Sheets XLSX exports often contain formatting for 1,000-2,000 rows
    across 20-30 columns. Walking the rectangular grid repeatedly caused the
    Test Series workbook to consume memory and remain in ``analyzing`` for
    many minutes. Reading the populated cell store keeps work proportional to
    real information.
    """

    grouped: dict[int, dict[int, Cell]] = defaultdict(dict)
    cell_store = getattr(worksheet, "_cells", {})
    candidates = cell_store.values() if cell_store else (
        cell for row in worksheet.iter_rows() for cell in row
    )
    for cell in candidates:
        if isinstance(cell, MergedCell):
            continue
        text = normalize_text(cell.value)
        link = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else None
        if not text and not link:
            continue
        grouped[cell.row][cell.column] = cell
    return dict(grouped)

def _cell_link(cell: Cell) -> str | None:
    if cell.hyperlink and cell.hyperlink.target:
        return str(cell.hyperlink.target).strip() or None
    text = normalize_text(cell.value)
    urls = find_urls(text) if text else []
    return urls[0] if urls else None


def profile_worksheet(worksheet: Worksheet, max_header_scan: int = 40) -> SheetSchemaProfile:
    rows = _populated_rows(worksheet)
    best_score = -1
    best_row: int | None = None
    best_columns: dict[str, int] = {}
    best_labels: dict[str, str] = {}
    scan_rows = min(max_header_scan, worksheet.max_row)
    for row_number in range(1, scan_rows + 1):
        columns: dict[str, int] = {}
        labels: dict[str, str] = {}
        for column, cell in rows.get(row_number, {}).items():
            original = normalize_text(cell.value)
            field_name = match_field(clean_header(original))
            if field_name and field_name not in columns:
                columns[field_name] = column
                labels[field_name] = original
        identity_count = len(IDENTITY_FIELDS.intersection(columns))
        file_count = len(FILE_FIELDS.intersection(columns))
        operation_count = len(TEST_OPERATION_FIELDS.intersection(columns))
        score = len(columns) + identity_count * 2 + file_count * 4 + operation_count * 2
        if score > best_score:
            best_score = score
            best_row = row_number if columns else None
            best_columns = columns
            best_labels = labels

    if not best_columns:
        return SheetSchemaProfile(
            missing_recommended_fields=["Title/Topic", "Subject or Class", "File/Link"],
            search_ready=False,
            confidence=0.0,
        )

    identity = IDENTITY_FIELDS.intersection(best_columns)
    files = FILE_FIELDS.intersection(best_columns)
    operations = TEST_OPERATION_FIELDS.intersection(best_columns)
    linked_rows = 0
    record_count = 0
    if best_row:
        important_columns = set(best_columns.values())
        for row_number, row_cells in rows.items():
            if row_number <= best_row:
                continue
            values = [
                normalize_text(row_cells[column].value)
                for column in important_columns
                if column in row_cells
            ]
            row_has_identity = sum(bool(value) for value in values) >= 2
            row_has_link = any(_cell_link(cell) for cell in row_cells.values())
            if row_has_identity:
                record_count += 1
            if row_has_link:
                linked_rows += 1

    # A new workbook may put links under a heading such as Remarks or Output.
    # Actual hyperlink evidence is stronger than the column label, so these tabs
    # remain searchable instead of silently losing hundreds of rows.
    search_ready = bool(identity and (files or operations or linked_rows))
    missing: list[str] = []
    if "topic" not in best_columns and "subject" not in best_columns:
        missing.append("Title/Topic or Subject")
    if "class_name" not in best_columns:
        missing.append("Class (recommended)")
    if "subject" not in best_columns:
        missing.append("Subject (recommended)")
    if not files and not linked_rows:
        missing.append("File/Link column")

    searchable_fields = sorted(best_columns)
    labels = dict(best_labels)
    if linked_rows and not files:
        searchable_fields.append("hyperlinks")
        labels["hyperlinks"] = f"Links found in {linked_rows} data rows"

    confidence = min(0.99, 0.35 + len(best_columns) * 0.055 + (0.15 if search_ready else 0.0) + (0.05 if linked_rows else 0.0))
    return SheetSchemaProfile(
        header_row=best_row,
        columns=best_columns,
        labels=labels,
        searchable_fields=searchable_fields,
        missing_recommended_fields=missing,
        search_ready=search_ready,
        record_count=record_count,
        confidence=confidence,
    )

def _cell_value_and_link(worksheet: Worksheet, row: int, column: int | None) -> tuple[str | None, str | None, str]:
    if not column:
        return None, None, f"A{row}"
    cell = worksheet.cell(row, column)
    title = normalize_text(cell.value) or None
    url = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else None
    if not url and title:
        urls = find_urls(title)
        url = urls[0] if urls else None
    return title, url, cell.coordinate


def _content_type(field_name: str, title: str | None, url: str | None) -> ContentType:
    if field_name == "pdf":
        return ContentType.PDF
    if field_name == "video":
        return ContentType.VIDEO
    if field_name == "presentation":
        return ContentType.PRESENTATION
    if field_name == "spreadsheet":
        return ContentType.SPREADSHEET
    if field_name == "document":
        return _classify_content_type(url or "", title or "document")
    return _classify_content_type(url or "", title or "")


_TEACHER_SUBJECTS: tuple[tuple[str, str], ...] = (
    ("ravi", "Physics"),
    ("kuldeep", "Mathematics"),
    ("apeksha", "Biology"),
    ("tapur", "Chemistry"),
    ("magnollia", "English"),
    ("magnolia", "English"),
    ("pranav", "Social Science"),
    ("kishore", "Mathematics"),
    ("lali", "Hindi"),
    ("pankaj", "Sanskrit"),
)

_SUBJECT_HINTS: tuple[tuple[tuple[str, ...], str], ...] = (
    (("physics", "electric charge", "ray optics", "mechanics", "thermodynamics"), "Physics"),
    (("chemistry", "molarity", "molality", "solution", "electrochemistry", "organic", "inorganic"), "Chemistry"),
    (("biology", "reproduction", "living world", "mammary", "botany", "zoology"), "Biology"),
    (("math", "mathematics", "relation and function", "calculus", "sets", "matrix", "derivative"), "Mathematics"),
    (("english", "grammar", "poem", "literature"), "English"),
    (("hindi", "हिंदी", "कबीर"), "Hindi"),
    (("sanskrit", "संस्कृत"), "Sanskrit"),
    (("political science", "civics"), "Political Science"),
    (("geography",), "Geography"),
    (("history",), "History"),
    (("economics",), "Economics"),
    (("accountancy", "accounts"), "Accountancy"),
    (("business studies", "bst"), "Business Studies"),
)


def _infer_subject(explicit: str | None, teacher: str | None, topic: str | None) -> str | None:
    if explicit:
        return canonical_subject(explicit) or explicit
    blob = f"{teacher or ''} {topic or ''}".casefold()
    for hints, subject in _SUBJECT_HINTS:
        if any(hint in blob for hint in hints):
            return subject
    teacher_text = (teacher or "").casefold()
    for token, subject in _TEACHER_SUBJECTS:
        if token in teacher_text:
            return subject
    return None


def _strip_urls(value: str | None) -> str | None:
    text = normalize_text(value)
    if not text:
        return None
    for url in find_urls(text):
        text = text.replace(url, " ")
    text = normalize_text(text)
    return text or None


def _row_candidates(
    worksheet: Worksheet,
    row: int,
    columns: dict[str, int],
    row_cells: dict[int, Cell] | None = None,
) -> list[tuple[str, str | None, str | None, str]]:
    """Return every useful file/link in a row, even under a misleading heading.

    The tuple is (field_name, title, url, coordinate). Recognised file columns
    can keep a filename without a URL. Any other column is included only when
    the cell contains a real hyperlink or URL.
    """
    candidates: list[tuple[str, str | None, str | None, str]] = []
    seen: set[tuple[str, str]] = set()
    recognised_columns = {column: field for field, column in columns.items() if field in FILE_FIELDS}
    for column, field_name in recognised_columns.items():
        title, url, coordinate = _cell_value_and_link(worksheet, row, column)
        if _usable_structured_file(field_name, title, url):
            key = (coordinate, url or normalize_text(title))
            seen.add(key)
            candidates.append((field_name, title, url, coordinate))

    physical_cells = (row_cells or {}).values() if row_cells is not None else (
        worksheet.cell(row, column) for column in range(1, worksheet.max_column + 1)
    )
    for cell in physical_cells:
        url = _cell_link(cell)
        if not url:
            continue
        title = normalize_text(cell.value) or None
        coordinate = cell.coordinate
        key = (coordinate, url)
        if key in seen:
            continue
        seen.add(key)
        candidates.append(("linked_cell", title, url, coordinate))
    return candidates


def _content_type_for_row(field_name: str, row_type: str | None, title: str | None, url: str | None) -> ContentType:
    inferred = _content_type(field_name, title, url)
    if inferred != ContentType.EXTERNAL_LINK:
        return inferred
    type_text = (row_type or "").casefold()
    if any(token in type_text for token in ("video", "short", "live", "youtube")):
        return ContentType.VIDEO
    if "pdf" in type_text:
        return ContentType.PDF
    if any(token in type_text for token in ("test", "doc", "paper")):
        return ContentType.DOCUMENT
    if any(token in type_text for token in ("slide", "ppt", "presentation")):
        return ContentType.PRESENTATION
    return inferred


def parse_structured_content_workbook(
    workbook_path: Path,
    *,
    source_id: str,
    workbook_name: str,
) -> list[ContentItem]:
    """Index explained rows from any newly connected Google Sheet.

    Heading aliases provide metadata, but actual hyperlinks are discovered in
    every physical column. This prevents a valid link column named Remarks,
    Output, Raw, Status, etc. from being ignored.
    """

    workbook_input = workbook_path if workbook_path.suffix.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"} else BytesIO(workbook_path.read_bytes())
    workbook = load_workbook(workbook_input, data_only=False, read_only=False)
    items: list[ContentItem] = []
    for worksheet in workbook.worksheets:
        populated_rows = _populated_rows(worksheet)
        profile = profile_worksheet(worksheet)
        if not profile.search_ready or not profile.header_row:
            continue
        columns = profile.columns
        for row in sorted(row_number for row_number in populated_rows if row_number > profile.header_row):
            row_cells = populated_rows[row]
            row_data = {
                field_name: normalize_text(row_cells[column].value) or None
                for field_name, column in columns.items()
                if column in row_cells
            }
            candidates = _row_candidates(worksheet, row, columns, row_cells)
            if not candidates:
                continue

            topic = row_data.get("topic") or row_data.get("chapter_number")
            teacher = row_data.get("teacher")
            explicit_class = row_data.get("class_name")
            class_name = (
                _detect_class(explicit_class or "")
                or explicit_class
                or _detect_class(worksheet.title)
                or _detect_class(workbook_name)
            )
            course = row_data.get("course")
            subject = _infer_subject(row_data.get("subject"), teacher, topic)
            stream = row_data.get("stream") or _detect_stream(" ".join(
                value for value in (course, class_name, subject, topic, worksheet.title, workbook_name) if value
            ))

            for field_name, cell_title, url, coordinate in candidates:
                clean_cell_title = _strip_urls(cell_title)
                # Topic/Title describes the row; a link cell often contains only
                # a URL or a generic value such as "Open".
                display_title = topic or clean_cell_title or subject or "Untitled learning item"
                if display_title.casefold() in _FILE_PLACEHOLDERS:
                    display_title = topic or subject or "Untitled learning item"
                content_type = _content_type_for_row(field_name, row_data.get("content_type"), cell_title, url)
                raw_search = " ".join(
                    value for value in (
                        workbook_name, worksheet.title, row_data.get("date"), row_data.get("content_type"),
                        course, stream, class_name, subject, row_data.get("chapter_number"), topic,
                        teacher, clean_cell_title, display_title, url, row_data.get("app_status"),
                        row_data.get("test_date"), row_data.get("result_dashboard"),
                        row_data.get("video_solution"), row_data.get("questions"),
                        row_data.get("duration"), row_data.get("marks"), row_data.get("remarks"),
                    ) if value
                )
                item_id = stable_hash("structured-content-v2", source_id, worksheet.title, row, coordinate, url or display_title)
                warnings = [] if url else ["The row is searchable, but a clickable file link is not added"]
                items.append(ContentItem(
                    id=item_id,
                    source_id=source_id,
                    workbook_name=workbook_name,
                    sheet_name=worksheet.title,
                    coordinate=coordinate,
                    source_row=row,
                    url=url or "",
                    title=display_title,
                    content_type=content_type,
                    stream=stream,
                    class_name=class_name,
                    batch=course,
                    subject=subject,
                    topic=topic or display_title,
                    teacher=teacher,
                    search_text=raw_search,
                    confidence=0.99 if url else 0.84,
                    warnings=warnings,
                    fingerprint=stable_hash(source_id, worksheet.title, row, coordinate, url or display_title),
                    record_kind="structured_link" if field_name == "linked_cell" else "structured",
                    status=row_data.get("app_status"),
                    test_number=row_data.get("chapter_number"),
                    test_date=row_data.get("test_date"),
                    result_dashboard=row_data.get("result_dashboard"),
                    video_solution_update=row_data.get("video_solution"),
                    questions=row_data.get("questions"),
                    duration=row_data.get("duration"),
                    marks=row_data.get("marks"),
                    remarks=row_data.get("remarks"),
                ))
    unique = {item.id: item for item in items}
    return sorted(unique.values(), key=lambda item: (item.sheet_name, item.source_row, item.coordinate, item.title))

